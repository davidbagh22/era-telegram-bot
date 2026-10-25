from datetime import date
from io import BytesIO
from types import SimpleNamespace as NS
from unittest import IsolatedAsyncioTestCase, TestCase
from unittest.mock import AsyncMock, Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfReader

from app.services.executive_pdf_service import build_executive_pdf


def fixture(project_count=35):
    wb = Workbook()
    ws = wb.active
    ws.title = '01 · Сводка'
    ws.append(['ЭРА'])
    ws.append(['Период'])
    ws.append([])
    ws.append(['Показатель', 'Значение', 'Описание', 'Источник'])
    ws.append(['Текущий состав', 66, 'Одобренные участники', 'User'])
    ws.append(['Новые за период', 8, 'Первые одобрения за выбранные даты', 'AuditLog'])
    ws.append(['События за период', 1, 'По календарной дате', 'Event'])
    ws.append(['Проекты за период', project_count + 1, 'Созданы или имеют подтверждённую работу', 'Project'])
    ws.append(['Подтверждённых посещений', 19, 'Подтверждённая явка', 'EventRegistration'])
    ws.append(['Verified task results', 4, 'Приняты по аудиту', 'AuditLog'])
    ws = wb.create_sheet('02 · Люди')
    ws.append(['Участник', 'Статус'])
    ws.append(['Имя <script> & текст', 'Одобрено'])
    ws.append(['=Имя участника', '=Название проекта'])
    ws = wb.create_sheet('03 · Пустой раздел')
    ws.append(['Запись', 'Статус'])
    ws.append(['Данных за выбранный период нет', ''])
    ws = wb.create_sheet('04 · Длинные записи')
    ws.append(['Запись', 'Описание'])
    ws.append(['Подробности', 'Очень длинный текст <без разметки> ' * 700])
    ws = wb.create_sheet('08 · Проекты')
    ws.append(['Проект', 'Текущий статус', 'Создан', 'Подтверждённых вкладов за период', 'Завершено этапов за период', 'Основание включения'])
    ws.append(['=Название проекта <script> & текст', 'in_progress', date(2026,10,1), 3, 2, 'Новый проект'])
    for i in range(project_count):
        ws.append([f'Проект {i} / Очень длинный текст ' + 'содержательная инициатива ' * 8, 'approved', date(2026,10,2), 0, 0, 'Новый проект'])
    ws = wb.create_sheet('07 · События')
    ws.append(['Событие', 'Дата', 'Статус', 'Регистраций', 'Посетили'])
    ws.append(['Встреча ЭРА', date(2026,10,20), 'completed', 24, 19])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report(project_count=35):
    return build_executive_pdf(fixture(project_count), period_label='01.10.2026–21.10.2026',
        health=NS(risks=['Сроки требуют внимания'], metrics=[], period_label='30 дней'),
        efficiency=NS(recommendations=[NS(title='Проверить сроки проектов', reason='Есть просроченные этапы', action='Обновить план работ')], period_label='7 дней'),
        analytics=NS(contacts=[NS(organization_name='Партнёр ЭРА', email='PRIVATE_EMAIL', phone='PRIVATE_PHONE', notes='PRIVATE_NOTE')],
                     goals=[NS(title='Цель организации', scope_name='ЭРА', month='2026-10', current_value=3, target_value=5, status='active')]),
        media=NS(published=12))


class ExecutivePDFTests(TestCase):
    def test_unicode_escaping_all_sections_and_long_pagination(self):
        content = report()
        self.assertTrue(content.startswith(b'%PDF-'))
        pdf = PdfReader(BytesIO(content))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        for expected in ['Текущий состав', '66', 'Приоритеты руководителя',
                         '=Название проекта <script> & текст', 'Проекты за период',
                         '01.10.2026', 'Очень длинный текст', 'Встреча ЭРА', 'Посетили: 19', 'Завершено этапов: 2',
                         'Сроки требуют внимания', 'Проверить сроки проектов', 'Есть просроченные этапы', 'Обновить план работ']:
            self.assertIn(expected, text)
        for secret in ['PRIVATE_EMAIL', 'PRIVATE_PHONE', 'PRIVATE_NOTE', 'Имя <script> & текст', '=Имя участника']:
            self.assertNotIn(secret, text)
        self.assertGreater(len(pdf.pages), 4)
        self.assertTrue(all('Для внутреннего использования' in page.extract_text() for page in pdf.pages))


class ExecutivePDFRouteTests(TestCase):
    def setUp(self):
        from app.api.v1 import admin_executive_export as api
        self.api = api
        app = FastAPI()
        app.include_router(api.router)
        app.dependency_overrides[api.get_session] = lambda: AsyncMock(add=Mock())
        app.dependency_overrides[api.get_settings] = lambda: NS(timezone='Asia/Yerevan', era_channel_id=None, general_chat_id=None)
        app.dependency_overrides[api.require_admin] = lambda: NS(id=1)
        self.app = app
        self.client = TestClient(app)

    def test_period_pdf_headers_and_shared_facts(self):
        api = self.api
        with patch.object(api, 'telegram_snapshot', AsyncMock(return_value={'start': '2026-10-01', 'end': '2026-10-21', 'sources': []})), \
             patch.object(api, 'build_analytics_payload', AsyncMock(return_value=NS())), \
             patch.object(api, 'build_efficiency_snapshot', AsyncMock(return_value=NS())), \
             patch.object(api, 'build_extended_organization_health', AsyncMock(return_value=NS())), \
             patch.object(api, 'media_service') as media_service, \
             patch.object(api, 'build_executive_workbook', AsyncMock(return_value=fixture())) as builder:
            media_service.can_manage_media = AsyncMock(return_value=False)
            media_service.analytics = AsyncMock()
            response = self.client.get('/admin/analytics/executive-report.pdf?period=custom&start_date=2026-10-01&end_date=2026-10-21')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/pdf')
        self.assertEqual(response.headers['cache-control'], 'no-store, private')
        self.assertIn('20261001-20261021.pdf', response.headers['content-disposition'])
        media_service.analytics.assert_not_awaited()
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertEqual(builder.call_args.kwargs['start_date'], date(2026, 10, 1))

    def test_invalid_period_does_not_build_report(self):
        with patch.object(self.api, 'build_analytics_payload', AsyncMock()) as builder:
            response = self.client.get('/admin/analytics/executive-report.pdf?period=custom&start_date=2026-10-21&end_date=2026-10-01')
        self.assertEqual(response.status_code, 422)
        builder.assert_not_awaited()
        self.assertEqual(self.client.get('/admin/analytics/executive-report.exe').status_code, 422)


class ExecutivePDFAuthorizationTests(IsolatedAsyncioTestCase):
    async def test_category_access_does_not_allow_export(self):
        from fastapi import HTTPException
        from app.api.v1.admin_executive_export import require_admin
        with patch('app.api.v1.admin_executive_export.has_dashboard_access', return_value=False):
            with self.assertRaises(HTTPException) as denied:
                await require_admin(NS(telegram_id=7), NS())
        self.assertEqual(denied.exception.status_code, 403)
