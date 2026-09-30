from io import BytesIO

from openpyxl import load_workbook

from app.commission_bot_analytics import build_analytics_workbook, should_interrupt_flow


def test_main_navigation_recovers_from_stale_wizard_state():
    for action in (
        "➕ Создать",
        "🎟 Мои заявки",
        "👤 Профиль",
        "📊 Аналитика",
        "💬 Обращения",
    ):
        assert should_interrupt_flow("b:description", True, action) is True


def test_incomplete_onboarding_is_not_silently_abandoned():
    assert should_interrupt_flow("onboard:country", False, "👤 Профиль") is False
    assert should_interrupt_flow("onboard:country", False, "📊 Аналитика") is False


def test_help_and_non_navigation_messages_do_not_cancel_flow():
    assert should_interrupt_flow("b:title", True, "❓ Помощь") is False
    assert should_interrupt_flow("support:reply:12", True, "обычный текст") is False
    assert should_interrupt_flow(None, True, "👤 Профиль") is False


def test_admin_analytics_workbook_contains_all_expected_sections():
    snapshot = {
        "summary": {
            "participants": 10,
            "profiles": 8,
            "countries": 3,
            "optins": 9,
            "publications": 5,
            "published": 4,
            "upcoming": 2,
            "registrations": 12,
            "unique_participants": 7,
            "attended": 6,
            "waitlist": 1,
            "sent": 40,
            "failed": 2,
            "targets": 4,
            "support_open": 1,
            "support_answered": 2,
            "support_closed": 3,
            "avg_first_response_min": 25,
        },
        "participants": [[1, "Test User", "@test", 22, "Армения", "Ереван", "Студент", "МГУ", None, None, None, "Культура", "Да", "Да", None, None]],
        "events": [[1, "Event", "event", "published", None, "offline", "Yerevan", 100, 12, 6, 1, 0, 40, 2, 50.0]],
        "registrations": [[1, "Event", "registered", None, 1, "Test User", "@test", 22, "Армения", "Ереван", "Студент", "МГУ", None, None, None, "Культура", {"custom": {"0": "Да"}}]],
        "distribution": [[1, "Event", "telegram", 42, 40, 2, 95.2]],
        "support": [[1, "Test User", "Армения", "Регистрация", "answered", None, None, "Admin", 25.0, 2]],
        "countries": [["Армения", 10, 8, 9, 12, 6]],
        "interests": [["Культура", 5]],
    }

    data = build_analytics_workbook(snapshot)
    workbook = load_workbook(BytesIO(data), read_only=False, data_only=False)

    assert workbook.sheetnames == [
        "Dashboard",
        "Participants",
        "Events",
        "Registrations",
        "Distribution",
        "Support",
        "Countries",
        "Interests",
    ]
    assert workbook["Dashboard"]["B4"].value == 10
    assert workbook["Participants"].freeze_panes == "A2"
    assert workbook["Participants"].auto_filter.ref is not None
    assert workbook["Events"]["I2"].value == 12
