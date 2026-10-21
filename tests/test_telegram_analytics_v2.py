"""Source facts, honest coverage, calendar boundaries, exports and access contracts."""
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace as NS
from unittest import IsolatedAsyncioTestCase, TestCase
from unittest.mock import AsyncMock, patch

from aiogram.types import Chat, Message, Update, User as TelegramUser
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.database.base import Base
from app.database.models import User
from app.database.telegram_analytics import TelegramAnalyticsSource, TelegramMessageFact
from app.services.telegram_analytics_service import observe_message, refresh_sources, report_window, telegram_snapshot, persist_coverage_gap, _coverage_gaps
from app.services.telegram_analytics_export import add_telegram_sheets

NOW = datetime(2026, 10, 21, 12, tzinfo=timezone.utc)
SETTINGS = NS(era_channel_id=-1001, general_chat_id=-1002, era_channel_username='era_fixture', timezone='Asia/Yerevan')


def message(id=1, *, channel=True, at=NOW, bot=False, **kwargs):
    if isinstance(kwargs.get('edit_date'), datetime):
        kwargs['edit_date'] = int(kwargs['edit_date'].timestamp())
    return Message(message_id=id, chat=Chat(id=-1001 if channel else -1002, type='channel' if channel else 'supergroup'),
                   date=at, from_user=TelegramUser(id=7, is_bot=bot, first_name='Fixture'), **({'text': 'NEVER_STORE_THIS_BODY'} | kwargs))


class TelegramFactsTests(IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        _coverage_gaps.clear()
        self.engine = create_async_engine('sqlite+aiosqlite:///:memory:')
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.factory = async_sessionmaker(self.engine, expire_on_commit=False)
        self.session = self.factory()
        self.start, self.end = report_window(days=7, now=NOW)

    async def asyncTearDown(self):
        await self.session.close()
        await self.engine.dispose()

    async def source(self, kind='channel', coverage=None, status='collecting', checked=NOW):
        row = TelegramAnalyticsSource(chat_id=SETTINGS.era_channel_id if kind == 'channel' else SETTINGS.general_chat_id,
            kind=kind, status=status, started_at=NOW-timedelta(days=100),
            coverage_since=(coverage or NOW-timedelta(days=100)).astimezone(timezone.utc), checked_at=checked)
        self.session.add(row)
        await self.session.flush()
        return row

    async def snapshot(self):
        return await telegram_snapshot(self.session, SETTINGS, start=self.start, end=self.end, now=NOW+timedelta(seconds=1))

    async def test_manual_bot_and_service_channel_posts_deduplicated(self):
        await self.source()
        for item in [message(), message(), message(edit_date=NOW+timedelta(minutes=1)), message(2, bot=True), message(3)]:
            await observe_message(self.session, item, kind='channel', observed_at=NOW)
        result = (await self.snapshot())['sources'][0]
        self.assertEqual(result['value'], 3)
        self.assertEqual(result['delta'], 3)
        self.assertEqual(result['data_state'], 'valid')
        self.assertEqual(await self.session.scalar(select(func.count()).select_from(TelegramMessageFact)), 3)
        row = await self.session.get(TelegramMessageFact, (-1001, 1))
        self.assertEqual(row.edited_at.replace(tzinfo=timezone.utc), NOW+timedelta(minutes=1))
        self.assertNotIn('text', TelegramMessageFact.__table__.columns)
        self.assertNotIn('user_id', TelegramMessageFact.__table__.columns)

    async def test_album_one_publication_across_local_midnight_and_edits(self):
        await self.source()
        at = datetime(2026, 10, 20, 19, 59, tzinfo=timezone.utc)
        for item in [message(1, at=at, media_group_id='album'), message(2, at=at+timedelta(minutes=2), media_group_id='album'), message(2, at=at+timedelta(minutes=2), media_group_id='album', edit_date=NOW)]:
            await observe_message(self.session, item, kind='channel', observed_at=NOW)
        source = (await self.snapshot())['sources'][0]
        self.assertEqual(source['value'], 1)
        daily = {row['date']: row['value'] for row in source['daily']}
        self.assertEqual(daily['2026-10-20'], 1)
        self.assertEqual(daily['2026-10-21'], 0)

    async def test_topics_admins_and_anonymous_admin_count_bots_service_forward_do_not(self):
        await self.source('general')
        for item in [message(1, channel=False, message_thread_id=7), message(2, channel=False, message_thread_id=8),
                     message(3, channel=False, bot=True, sender_chat=Chat(id=-1002, type='supergroup')),
                     message(7, channel=False, bot=True, sender_chat=Chat(id=-1009, type='channel')),
                     message(4, channel=False, bot=True), message(5, channel=False, is_automatic_forward=True),
                     message(6, channel=False, text=None, new_chat_title='Service')]:
            await observe_message(self.session, item, kind='general', observed_at=NOW)
        self.assertEqual((await self.snapshot())['sources'][1]['value'], 4)

    async def test_partial_coverage_is_not_zero_or_comparable(self):
        await self.source(coverage=NOW-timedelta(hours=1))
        await observe_message(self.session, message(), kind='channel', observed_at=NOW)
        row = (await self.snapshot())['sources'][0]
        self.assertEqual(row['data_state'], 'insufficient')
        self.assertIsNone(row['value'])
        self.assertIsNone(row['previous_value'])
        self.assertEqual(row['observed_value'], 1)
        self.assertTrue(all(day['value'] is None for day in row['daily']))

    async def test_zero_requires_full_coverage_and_previous_requires_two_periods(self):
        await self.source(coverage=self.start)
        row = (await self.snapshot())['sources'][0]
        self.assertEqual(row['value'], 0)
        self.assertIsNone(row['previous_value'])
        self.assertIsNone(row['delta'])

    async def test_unconfigured_limited_stale_sources_are_not_false_zeroes(self):
        self.assertIsNone((await self.snapshot())['sources'][0]['value'])
        row = await self.source(status='limited')
        self.assertEqual((await self.snapshot())['sources'][0]['data_state'], 'unavailable')
        row.status = 'collecting'
        row.checked_at = NOW-timedelta(minutes=16)
        await self.session.flush()
        self.assertEqual((await self.snapshot())['sources'][0]['data_state'], 'stale')
        self.assertIsNone((await self.snapshot())['sources'][0]['value'])

    async def test_permission_loss_and_gap_reset_continuous_coverage(self):
        row = await self.source()
        bot = NS(id=99, get_chat_member=AsyncMock(return_value=NS(status='member')), get_me=AsyncMock(return_value=NS(can_read_all_group_messages=False)))
        await refresh_sources(self.session, bot, SETTINGS, now=NOW)
        self.assertEqual(row.status, 'unavailable')
        self.assertIsNone(row.coverage_since)
        bot.get_chat_member.return_value = NS(status='administrator')
        await refresh_sources(self.session, bot, SETTINGS, now=NOW+timedelta(minutes=1))
        self.assertEqual(row.coverage_since, NOW+timedelta(minutes=1))
        await refresh_sources(self.session, bot, SETTINGS, now=NOW+timedelta(minutes=21))
        self.assertEqual(row.coverage_since, NOW+timedelta(minutes=21))

    async def test_gap_invalidation_is_visible_after_worker_state_is_lost(self):
        await self.source()
        await self.session.commit()
        async with self.factory() as writer:
            await persist_coverage_gap(writer, SETTINGS.era_channel_id)
            await writer.commit()
        _coverage_gaps.clear()  # Another worker has no in-memory signal.
        async with self.factory() as reader:
            result = await telegram_snapshot(reader, SETTINGS, start=self.start, end=self.end, now=NOW)
        self.assertEqual(result['sources'][0]['connection_status'], 'unverified')
        self.assertIsNone(result['sources'][0]['coverage_since'])
        self.assertIsNone(result['sources'][0]['value'])

    async def test_pending_gap_cannot_certify_underreported_counts(self):
        from app.services.telegram_analytics_service import mark_coverage_gap
        await self.source()
        mark_coverage_gap(SETTINGS.era_channel_id)
        source = (await self.snapshot())['sources'][0]
        self.assertEqual(source['data_state'], 'insufficient')
        self.assertIsNone(source['value'])
        self.assertTrue(all(day['value'] is None for day in source['daily']))

    async def test_probe_does_not_clear_pending_gap_before_commit(self):
        from app.services.telegram_analytics_service import mark_coverage_gap, confirm_sources_refresh
        await self.source()
        mark_coverage_gap(SETTINGS.era_channel_id)
        bot = NS(id=99, get_chat_member=AsyncMock(return_value=NS(status='administrator')))
        acknowledgements = await refresh_sources(self.session, bot, SETTINGS, now=NOW)
        self.assertIn(SETTINGS.era_channel_id, _coverage_gaps)
        await self.session.commit()
        confirm_sources_refresh(acknowledgements)
        self.assertNotIn(SETTINGS.era_channel_id, _coverage_gaps)
        # A later gap must survive acknowledgement of an older probe.
        acknowledgements = await refresh_sources(self.session, bot, SETTINGS, now=NOW+timedelta(minutes=1))
        mark_coverage_gap(SETTINGS.era_channel_id)
        await self.session.commit()
        confirm_sources_refresh(acknowledgements)
        self.assertIn(SETTINGS.era_channel_id, _coverage_gaps)

    async def test_real_roster_and_all_profiles_have_distinct_definitions(self):
        from app.services.admin_analytics_v2_service import build_v2_snapshot
        self.session.add_all([User(telegram_id=i, first_name='Fixture', application_status='approved' if i < 68 else 'pending') for i in range(1,69)])
        await self.session.flush()
        snapshot = await build_v2_snapshot(self.session, SETTINGS, start=self.start, end=self.end)
        metrics = {row['id']: row for row in snapshot['metrics']}
        self.assertEqual(metrics['roster']['value'], 67)
        self.assertEqual(metrics['records']['value'], 68)
        self.assertIsNone(metrics['roster']['previous_value'])
        self.assertEqual(metrics['verified_active']['denominator'], 67)

    async def test_empty_health_has_no_zero_denominator_percentage(self):
        from app.services.organization_health_extended_service import build_extended_organization_health
        snapshot = await build_extended_organization_health(self.session)
        self.assertEqual(len(snapshot.metrics), 59)
        active = next(metric for metric in snapshot.metrics if metric.key == 'active_30d')
        self.assertNotIn('0.0%', active.display)
        self.assertIsNone(active.score)
        for metric in snapshot.metrics:
            if metric.unit != 'count' and not metric.denominator:
                self.assertIsNone(metric.value, metric.key)
                self.assertEqual(metric.data_state, 'unavailable', metric.key)

    async def test_middleware_observes_without_consuming_unhandled_message(self):
        from app.middlewares.telegram_analytics import TelegramAnalyticsMiddleware
        handler = AsyncMock(return_value='business result')
        result = await TelegramAnalyticsMiddleware(self.factory, SETTINGS)(handler, Update(update_id=1, message=message(channel=False)), {})
        self.assertEqual(result, 'business result')
        handler.assert_awaited_once()
        self.assertEqual(await self.session.scalar(select(func.count()).select_from(TelegramMessageFact)), 1)

    async def test_export_keeps_null_and_observed_counts_and_same_days(self):
        await self.source(coverage=NOW-timedelta(hours=1))
        await observe_message(self.session, message(), kind='channel', observed_at=NOW)
        snapshot = await self.snapshot()
        wb=Workbook(); buffer=BytesIO(); wb.save(buffer)
        exported = load_workbook(BytesIO(add_telegram_sheets(buffer.getvalue(), snapshot)))
        rows = list(exported['Telegram · источники'].values)
        self.assertIn('insufficient', rows[1])
        self.assertIn(1, rows[1])
        self.assertIn(None, rows[1])
        self.assertEqual(len(list(exported['Telegram · по дням'].values)), 15)
        self.assertNotIn('NEVER_STORE_THIS_BODY', str(rows))
        from app.services.executive_pdf_service import build_executive_pdf
        from pypdf import PdfReader
        pdf = build_executive_pdf(add_telegram_sheets(buffer.getvalue(), snapshot), period_label='Fixture',
            health=NS(), efficiency=NS(), telegram=snapshot)
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
        self.assertIn('Канал и общий чат', text)
        self.assertIn('Неполная история', text)
        self.assertIn('1 учтено', text)
        self.assertNotIn('NEVER_STORE_THIS_BODY', text)


class WindowTests(TestCase):
    def test_local_dates_and_supported_periods(self):
        for days in (7, 30, 90):
            start, end = report_window(days=days, now=NOW)
            self.assertEqual((end-start).days, days)
            self.assertEqual(end.date(), date(2026,10,22))
            self.assertEqual(start.utcoffset(), timedelta(hours=4))
        for kwargs in [dict(start_date=date(2026,1,1)), dict(start_date=date(2026,10,22), end_date=date(2026,10,23)), dict(start_date=date(2025,1,1), end_date=date(2026,10,21))]:
            with self.assertRaises(ValueError):
                report_window(now=NOW, **kwargs)


class AnalyticsAccessTests(IsolatedAsyncioTestCase):
    async def test_actual_query_string_periods_and_custom_dates_are_accepted(self):
        from app.api.v1 import admin_analytics_v2 as api
        app=FastAPI(); app.include_router(api.router)
        app.dependency_overrides[api.require_admin]=lambda: NS(id=1)
        app.dependency_overrides[api.get_session]=lambda: None
        app.dependency_overrides[api.get_settings]=lambda: SETTINGS
        client=TestClient(app)
        with patch.object(api, 'build_v2_snapshot', AsyncMock(return_value={'ok': True})) as build:
            for days in (7,30,90):
                response=client.get(f'/admin/analytics/v2?days={days}')
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual((build.call_args.kwargs['end']-build.call_args.kwargs['start']).days,days)
            self.assertEqual(client.get('/admin/analytics/v2?days=8').status_code,422)
            response=client.get('/admin/analytics/v2?days=30&start_date=2026-01-01&end_date=2026-10-21')
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(build.call_args.kwargs['start'].date(),date(2026,1,1))

    async def test_category_grant_does_not_open_any_v2_endpoint(self):
        from app.api.v1 import admin_analytics_v2 as api
        app=FastAPI(); app.include_router(api.router)
        async def denied():
            raise HTTPException(403, 'admin_access_required')
        app.dependency_overrides[api.require_admin]=denied
        app.dependency_overrides[api.get_session]=lambda: None
        app.dependency_overrides[api.get_settings]=lambda: SETTINGS
        app.dependency_overrides[api.get_bot]=lambda: None
        client=TestClient(app)
        for method,url in [('get','/admin/analytics/v2'),('get','/admin/analytics/v2/reports'),('post','/admin/analytics/v2/sources/check')]:
            self.assertEqual(getattr(client,method)(url).status_code,403)
