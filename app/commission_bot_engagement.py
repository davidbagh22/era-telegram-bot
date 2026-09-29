from __future__ import annotations

import asyncio
import html
import io
import json
import logging
from datetime import datetime, timedelta, timezone
from typing import Any
from zoneinfo import ZoneInfo

from aiogram import F
from aiogram.dispatcher.event.bases import SkipHandler
from aiogram.enums import ChatType
from aiogram.exceptions import TelegramNetworkError, TelegramRetryAfter, TelegramServerError
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Font

from app.commission_bot import _kb
from app.commission_bot_analytics import CommissionBotAnalytics
from app.commission_bot_community import INTEREST_LABELS, STATUS_LABELS, _json_list
from app.commission_bot_region import REGION_COUNTRIES, REGION_LABELS, WORLD_COUNTRIES, _country_label

log = logging.getLogger(__name__)

ENGAGEMENT_DDL = r"""
ALTER TABLE users ADD COLUMN IF NOT EXISTS last_activity_at TIMESTAMPTZ NOT NULL DEFAULT NOW();
ALTER TABLE users ADD COLUMN IF NOT EXISTS monthly_survey_opt_in BOOLEAN NOT NULL DEFAULT TRUE;
ALTER TABLE users ADD COLUMN IF NOT EXISTS activity_level TEXT NOT NULL DEFAULT 'new';
ALTER TABLE targets ADD COLUMN IF NOT EXISTS geography_scope TEXT NOT NULL DEFAULT 'country';
ALTER TABLE targets ADD COLUMN IF NOT EXISTS region_code TEXT;
ALTER TABLE targets ADD COLUMN IF NOT EXISTS purpose TEXT NOT NULL DEFAULT 'general';
ALTER TABLE targets ADD COLUMN IF NOT EXISTS tags JSONB NOT NULL DEFAULT '[]'::jsonb;

CREATE TABLE IF NOT EXISTS user_tags (
  telegram_id BIGINT NOT NULL REFERENCES users(telegram_id) ON DELETE CASCADE,
  tag TEXT NOT NULL,
  source TEXT NOT NULL DEFAULT 'manual',
  created_by BIGINT,
  created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
  PRIMARY KEY (telegram_id, tag)
);

CREATE TABLE IF NOT EXISTS survey_runs (
  id BIGSERIAL PRIMARY KEY,
  month_key TEXT NOT NULL UNIQUE,
  template_key TEXT NOT NULL,
  status TEXT NOT NULL DEFAULT 'sending',
  target_count INTEGER NOT NULL DEFAULT 0,
  sent_count INTEGER NOT NULL DEFAULT 0,
  completed_count INTEGER NOT NULL DEFAULT 0,
  scheduled_for TIMESTAMPTZ,
  started_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
  completed_at TIMESTAMPTZ,
  created_by BIGINT
);
CREATE TABLE IF NOT EXISTS survey_recipients (
  run_id BIGINT NOT NULL REFERENCES survey_runs(id) ON DELETE CASCADE,
  telegram_id BIGINT NOT NULL REFERENCES users(telegram_id) ON DELETE CASCADE,
  status TEXT NOT NULL DEFAULT 'pending',
  attempts INTEGER NOT NULL DEFAULT 0,
  error TEXT,
  invited_at TIMESTAMPTZ,
  completed_at TIMESTAMPTZ,
  PRIMARY KEY (run_id, telegram_id)
);
CREATE INDEX IF NOT EXISTS idx_survey_recipients_due ON survey_recipients(run_id,status,attempts);
CREATE TABLE IF NOT EXISTS survey_progress (
  run_id BIGINT NOT NULL REFERENCES survey_runs(id) ON DELETE CASCADE,
  telegram_id BIGINT NOT NULL REFERENCES users(telegram_id) ON DELETE CASCADE,
  question_index INTEGER NOT NULL DEFAULT 0,
  answers JSONB NOT NULL DEFAULT '{}'::jsonb,
  updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
  PRIMARY KEY (run_id, telegram_id)
);
CREATE TABLE IF NOT EXISTS survey_responses (
  run_id BIGINT NOT NULL REFERENCES survey_runs(id) ON DELETE CASCADE,
  telegram_id BIGINT NOT NULL REFERENCES users(telegram_id) ON DELETE CASCADE,
  answers JSONB NOT NULL DEFAULT '{}'::jsonb,
  completed_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
  PRIMARY KEY (run_id, telegram_id)
);

CREATE TABLE IF NOT EXISTS segment_campaigns (
  id BIGSERIAL PRIMARY KEY,
  created_by BIGINT NOT NULL REFERENCES users(telegram_id),
  body TEXT NOT NULL,
  segment JSONB NOT NULL DEFAULT '{}'::jsonb,
  channels JSONB NOT NULL DEFAULT '{}'::jsonb,
  status TEXT NOT NULL DEFAULT 'draft',
  estimated_users INTEGER NOT NULL DEFAULT 0,
  estimated_targets INTEGER NOT NULL DEFAULT 0,
  sent_users INTEGER NOT NULL DEFAULT 0,
  sent_targets INTEGER NOT NULL DEFAULT 0,
  failed INTEGER NOT NULL DEFAULT 0,
  created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
  sent_at TIMESTAMPTZ
);
CREATE TABLE IF NOT EXISTS segment_campaign_deliveries (
  campaign_id BIGINT NOT NULL REFERENCES segment_campaigns(id) ON DELETE CASCADE,
  channel TEXT NOT NULL,
  recipient_key TEXT NOT NULL,
  status TEXT NOT NULL DEFAULT 'pending',
  attempts INTEGER NOT NULL DEFAULT 0,
  error TEXT,
  sent_at TIMESTAMPTZ,
  PRIMARY KEY (campaign_id, channel, recipient_key)
);
CREATE INDEX IF NOT EXISTS idx_segment_delivery_status ON segment_campaign_deliveries(campaign_id,status,attempts);
"""

SURVEY_TEMPLATES: list[dict[str, Any]] = [
    {
        "key": "belonging",
        "title": "Мы и сообщество",
        "intro": "Начнём с главного: чувствуешь ли ты, что это действительно твоё сообщество?",
        "questions": [
            "Я чувствую себя частью движения и понимаю, зачем мы вместе.",
            "У меня есть желание участвовать активнее в ближайшие месяцы.",
            "В сообществе легко знакомиться и находить единомышленников.",
            "Я готов(а) рекомендовать участие друзьям и коллегам.",
        ],
        "comment": "Что помогло бы тебе сильнее почувствовать себя частью сообщества?",
    },
    {
        "key": "communication",
        "title": "Коммуникация без шума",
        "intro": "Проверяем, доходит ли важное вовремя и не теряется ли оно среди сообщений.",
        "questions": [
            "Я понимаю, где искать актуальную информацию.",
            "Сообщения движения приходят вовремя.",
            "Тексты понятные, короткие и помогают принять решение.",
            "Количество уведомлений для меня комфортное.",
        ],
        "comment": "Что в нашей коммуникации стоит изменить первым?",
    },
    {
        "key": "opportunities",
        "title": "Ценность возможностей",
        "intro": "Хотим понять, действительно ли возможности полезны, а не просто красиво выглядят в ленте.",
        "questions": [
            "Я регулярно нахожу здесь возможности, которые подходят мне.",
            "Участие даёт практическую пользу для учёбы, карьеры или проектов.",
            "Мне хватает разнообразия форматов и тем.",
            "Я понимаю, как воспользоваться опубликованными возможностями.",
        ],
        "comment": "Каких возможностей тебе сейчас не хватает?",
    },
    {
        "key": "online_events",
        "title": "Онлайн-события",
        "intro": "Смотрим на качество наших онлайн-встреч глазами участника.",
        "questions": [
            "Темы онлайн-событий мне интересны.",
            "Формат и продолжительность встреч удобны.",
            "После встреч у меня остаются конкретные идеи или полезные контакты.",
            "Я хотел(а) бы участвовать в следующих онлайн-событиях.",
        ],
        "comment": "Что сделало бы следующую онлайн-встречу заметно лучше?",
    },
    {
        "key": "trust",
        "title": "Доверие и обратная связь",
        "intro": "Сильное сообщество умеет не только говорить, но и слышать.",
        "questions": [
            "Я доверяю команде и понимаю логику принимаемых решений.",
            "Я могу безопасно высказать идею, вопрос или критику.",
            "Обратная связь участников влияет на реальные изменения.",
            "Команда отвечает на обращения достаточно быстро.",
        ],
        "comment": "Что повысило бы твоё доверие к работе команды?",
    },
    {
        "key": "growth",
        "title": "Личный рост",
        "intro": "Проверяем, помогает ли участие становиться сильнее, а не просто занимать время.",
        "questions": [
            "За последние месяцы я получил(а) новые знания или навыки.",
            "Участие помогает мне увереннее проявлять инициативу.",
            "Я вижу возможности для своего дальнейшего развития внутри сообщества.",
            "Полученный опыт можно применить за пределами движения.",
        ],
        "comment": "Какой навык или опыт ты хотел(а) бы получить дальше?",
    },
    {
        "key": "climate",
        "title": "Атмосфера сообщества",
        "intro": "Важно, чтобы активность строилась на уважении, открытости и нормальном человеческом общении.",
        "questions": [
            "В сообществе ко мне относятся с уважением.",
            "Новые участники могут быстро включиться в работу.",
            "Разные мнения можно обсуждать спокойно и конструктивно.",
            "Мне комфортно обращаться к другим участникам и команде.",
        ],
        "comment": "Что поможет сделать атмосферу ещё лучше?",
    },
    {
        "key": "networking",
        "title": "Связи без границ",
        "intro": "Измеряем не количество контактов, а то, появляются ли реальные связи между людьми и странами.",
        "questions": [
            "Благодаря движению я познакомился(лась) с людьми из других городов или стран.",
            "Новые знакомства перерастают в совместные идеи или проекты.",
            "Мне достаточно международных форматов и обменов.",
            "Я знаю, к кому обратиться, если ищу партнёра для идеи.",
        ],
        "comment": "С кем или с какой страной тебе хотелось бы больше взаимодействовать?",
    },
    {
        "key": "agency",
        "title": "Инициатива и проекты",
        "intro": "Хорошее движение не только приглашает участвовать — оно помогает запускать своё.",
        "questions": [
            "Я понимаю, как предложить собственную инициативу.",
            "Команда помогает довести хорошую идею до реализации.",
            "У меня достаточно свободы брать ответственность и пробовать новое.",
            "Я вижу понятный путь от участника к более активной роли.",
        ],
        "comment": "Что мешает тебе запускать больше собственных инициатив?",
    },
    {
        "key": "impact",
        "title": "Результат и следующий шаг",
        "intro": "Финальный ракурс: что движение реально меняет и куда ему идти дальше.",
        "questions": [
            "Движение создаёт заметную пользу для молодых людей.",
            "Я вижу реальные результаты, а не только количество мероприятий.",
            "За последний год движение стало сильнее и полезнее.",
            "Я хочу оставаться частью сообщества в следующем году.",
        ],
        "comment": "Если в следующем году можно улучшить только одну вещь — что это должно быть?",
    },
]
SURVEY_BY_KEY = {item["key"]: item for item in SURVEY_TEMPLATES}
SURVEY_BASE_YEAR = 2026
SURVEY_BASE_MONTH = 10
SURVEY_TZ = ZoneInfo("Asia/Yerevan")
SURVEY_DAY = 5
SURVEY_HOUR = 19

AGE_BUCKETS = {
    "u18": (0, 17, "до 18"),
    "18_25": (18, 25, "18–25"),
    "26_35": (26, 35, "26–35"),
    "36p": (36, 200, "36+"),
}
TARGET_PURPOSES = {
    "general": "Общий",
    "events": "Мероприятия",
    "opportunities": "Возможности",
    "education": "Образование",
    "announcements": "Объявления",
}
TARGET_TAGS = {
    "youth": "Молодёжь",
    "students": "Студенты",
    "leaders": "Лидеры",
    "volunteers": "Волонтёры",
    "media": "Медиа",
    "partners": "Партнёры",
}


def survey_template_for_month(year: int, month: int) -> dict[str, Any]:
    delta = (year * 12 + month) - (SURVEY_BASE_YEAR * 12 + SURVEY_BASE_MONTH)
    return SURVEY_TEMPLATES[delta % len(SURVEY_TEMPLATES)]


def country_region(code: str | None) -> str | None:
    if not code:
        return None
    for region, codes in REGION_COUNTRIES.items():
        if code in codes:
            return region
    return None


def _dict(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            decoded = json.loads(value)
            return decoded if isinstance(decoded, dict) else {}
        except Exception:
            return {}
    return {}


def user_matches_segment(user: dict[str, Any], segment: dict[str, Any]) -> bool:
    countries = set(segment.get("countries") or [])
    regions = set(segment.get("regions") or [])
    if countries and user.get("country_code") not in countries:
        return False
    if regions and country_region(user.get("country_code")) not in regions:
        return False
    age_buckets = set(segment.get("age_buckets") or [])
    if age_buckets:
        age = user.get("age")
        if age is None:
            return False
        if not any(lo <= int(age) <= hi for key, (lo, hi, _) in AGE_BUCKETS.items() if key in age_buckets):
            return False
    statuses = set(segment.get("statuses") or [])
    if statuses and user.get("participant_status") not in statuses:
        return False
    interests = set(segment.get("interests") or [])
    if interests and not interests.intersection(set(_json_list(user.get("interests")))):
        return False
    required_tags = set(segment.get("tags") or [])
    if required_tags and not required_tags.intersection(set(user.get("tags") or [])):
        return False
    if segment.get("active_only") and not bool(user.get("is_active")):
        return False
    return True


def target_matches_segment(target: dict[str, Any], segment: dict[str, Any]) -> bool:
    countries = set(segment.get("countries") or [])
    regions = set(segment.get("regions") or [])
    if not countries and not regions:
        return True
    scope = target.get("geography_scope") or "country"
    if scope == "global":
        return False
    if countries:
        return scope == "country" and target.get("country_code") in countries
    if regions:
        if scope == "region":
            return target.get("region_code") in regions
        return scope == "country" and country_region(target.get("country_code")) in regions
    return False


class CommissionBotEngagement(CommissionBotAnalytics):
    async def init_db(self) -> None:
        await super().init_db()
        await self.execute(ENGAGEMENT_DDL)

    async def ensure_user(self, u):
        row = await super().ensure_user(u)
        await self.execute("UPDATE users SET last_activity_at=NOW() WHERE telegram_id=$1", row["telegram_id"])
        return await self.fetchrow("SELECT * FROM users WHERE telegram_id=$1", row["telegram_id"])

    async def _show_management(self, chat_id: int, user) -> None:
        if user["role"] not in {"admin", "owner"}:
            return await self.bot.send_message(chat_id, "Этот раздел доступен администратору и владельцу.")
        rows = [
            [("🎯 Рассылки и сегменты", "seg:menu"), ("📣 Чаты и каналы", "tg:list")],
            [("📈 Эффективность", "svadm:menu"), ("📊 Аналитика", "ux:analytics")],
            [("💬 Обращения", "support:staff:list")],
        ]
        if user["role"] == "owner":
            rows.append([("👥 Команда", "ux:team")])
        rows.extend([
            [("🔔 Уведомления", "ux:notifications"), ("👤 Профиль", "ux:profile")],
            [("❓ Как пользоваться", "ux:help")],
            [("⬅️ Главное меню", "menu")],
        ])
        await self.bot.send_message(
            chat_id,
            "<b>⚙️ Управление</b>\n\nРассылки, сегменты, каналы, эффективность движения и вся служебная работа — в одном месте.",
            reply_markup=_kb(rows),
        )

    async def _active_audience_rows(self):
        rows = await self.fetch(
            """SELECT u.*,
            COALESCE((SELECT jsonb_agg(ut.tag) FROM user_tags ut WHERE ut.telegram_id=u.telegram_id),'[]'::jsonb) tag_list,
            EXISTS(SELECT 1 FROM registrations r WHERE r.telegram_id=u.telegram_id AND r.registered_at>=NOW()-INTERVAL '180 days' AND r.status<>'cancelled') recent_registration
            FROM users u
            WHERE u.onboarding_complete=TRUE AND u.telegram_opt_in=TRUE AND u.monthly_survey_opt_in=TRUE"""
        )
        result = []
        now = datetime.now(timezone.utc)
        for row in rows:
            d = dict(row)
            tags = _json_list(d.pop("tag_list", []))
            d["tags"] = tags
            recent_activity = d.get("last_activity_at") and d["last_activity_at"] >= now - timedelta(days=60)
            d["is_active"] = bool(recent_activity or d.get("recent_registration") or d.get("role") in {"editor", "admin", "owner"} or "activist" in tags)
            if d["is_active"]:
                result.append(d)
        return result

    async def _segment_candidates(self):
        rows = await self.fetch(
            """SELECT u.*,
            COALESCE((SELECT jsonb_agg(ut.tag) FROM user_tags ut WHERE ut.telegram_id=u.telegram_id),'[]'::jsonb) tag_list,
            EXISTS(SELECT 1 FROM registrations r WHERE r.telegram_id=u.telegram_id AND r.registered_at>=NOW()-INTERVAL '180 days' AND r.status<>'cancelled') recent_registration
            FROM users u WHERE u.onboarding_complete=TRUE AND u.telegram_opt_in=TRUE"""
        )
        now = datetime.now(timezone.utc)
        out = []
        for row in rows:
            d = dict(row)
            d["tags"] = _json_list(d.pop("tag_list", []))
            d["is_active"] = bool(
                (d.get("last_activity_at") and d["last_activity_at"] >= now - timedelta(days=60))
                or d.get("recent_registration")
                or d.get("role") in {"editor", "admin", "owner"}
                or "activist" in d["tags"]
            )
            out.append(d)
        return out

    async def _estimate_segment(self, segment: dict[str, Any], channels: dict[str, bool]) -> tuple[int, int]:
        users_count = 0
        targets_count = 0
        if channels.get("dm"):
            users_count = sum(1 for u in await self._segment_candidates() if user_matches_segment(u, segment))
        if channels.get("targets"):
            targets = await self.fetch("SELECT * FROM targets WHERE status='approved' AND bot_can_post=TRUE")
            targets_count = sum(1 for row in targets if target_matches_segment(dict(row), segment))
        return users_count, targets_count

    def _segment_summary(self, p: dict[str, Any]) -> str:
        seg = p.get("segment", {})
        channels = p.get("channels", {"dm": True, "targets": False})
        geo = "все страны"
        if seg.get("countries"):
            geo = ", ".join(_country_label(c) for c in seg["countries"] if c in WORLD_COUNTRIES)
        elif seg.get("regions"):
            geo = ", ".join(REGION_LABELS.get(r, r) for r in seg["regions"])
        filters = []
        if seg.get("age_buckets"):
            filters.append("возраст: " + ", ".join(AGE_BUCKETS[x][2] for x in seg["age_buckets"] if x in AGE_BUCKETS))
        if seg.get("statuses"):
            filters.append("статус: " + ", ".join(STATUS_LABELS.get(x, x) for x in seg["statuses"]))
        if seg.get("interests"):
            filters.append("интересы: " + ", ".join(INTEREST_LABELS.get(x, x) for x in seg["interests"]))
        if seg.get("active_only"):
            filters.append("только активные")
        delivery = []
        if channels.get("dm"):
            delivery.append("личные сообщения")
        if channels.get("targets"):
            delivery.append("чаты/каналы")
        return "🌍 " + geo + "\n🎯 " + ("; ".join(filters) if filters else "без дополнительных фильтров") + "\n📨 " + ", ".join(delivery)

    async def _show_segment_builder(self, chat_id: int, uid: int, p: dict[str, Any]) -> None:
        seg = p.setdefault("segment", {})
        channels = p.setdefault("channels", {"dm": True, "targets": False})
        users_count, targets_count = await self._estimate_segment(seg, channels)
        p["estimate"] = {"users": users_count, "targets": targets_count}
        await self.state_set(uid, "segment:builder", p)
        rows = [
            [("🌍 География", "seg:geo"), ("🎂 Возраст", "seg:age")],
            [("👤 Статус", "seg:status"), ("✨ Интересы", "seg:interest")],
            [("✅ Только активные" if seg.get("active_only") else "▫️ Только активные", "seg:active")],
            [("✅ Личные" if channels.get("dm") else "▫️ Личные", "seg:ch:dm"), ("✅ Чаты" if channels.get("targets") else "▫️ Чаты", "seg:ch:targets")],
            [(f"👥 Проверить аудиторию · {users_count + targets_count}", "seg:estimate")],
            [("✍️ Написать сообщение", "seg:compose")],
            [("✖️ Отмена", "seg:cancel")],
        ]
        await self.bot.send_message(
            chat_id,
            "<b>🎯 Новая точная рассылка</b>\n\n" + self._segment_summary(p) + f"\n\nСейчас получателей: <b>{users_count}</b> участников и <b>{targets_count}</b> чатов/каналов.\n\nСначала настрой аудиторию, потом напиши сообщение — перед отправкой будет ещё один контрольный экран.",
            reply_markup=_kb(rows),
        )

    async def _send_text_retry(self, chat_id: int, text: str):
        last: Exception | None = None
        for attempt in range(1, 4):
            try:
                msg = await self.bot.send_message(chat_id, html.escape(text), disable_web_page_preview=True)
                return msg, attempt, None
            except TelegramRetryAfter as exc:
                last = exc
                await asyncio.sleep(min(float(exc.retry_after) + 0.2, 30))
            except (TelegramNetworkError, TelegramServerError) as exc:
                last = exc
                await asyncio.sleep(attempt * 1.5)
            except Exception as exc:
                last = exc
                break
        return None, 3, last

    async def _deliver_segment_campaign(self, campaign_id: int) -> None:
        campaign = await self.fetchrow("SELECT * FROM segment_campaigns WHERE id=$1", campaign_id)
        if not campaign:
            return
        segment = _dict(campaign["segment"])
        channels = _dict(campaign["channels"])
        body = campaign["body"]
        if channels.get("dm"):
            for u in await self._segment_candidates():
                if not user_matches_segment(u, segment):
                    continue
                await self.execute(
                    "INSERT INTO segment_campaign_deliveries(campaign_id,channel,recipient_key) VALUES($1,'dm',$2) ON CONFLICT DO NOTHING",
                    campaign_id, str(u["telegram_id"]),
                )
        if channels.get("targets"):
            rows = await self.fetch("SELECT * FROM targets WHERE status='approved' AND bot_can_post=TRUE")
            for target in rows:
                if not target_matches_segment(dict(target), segment):
                    continue
                await self.execute(
                    "INSERT INTO segment_campaign_deliveries(campaign_id,channel,recipient_key) VALUES($1,'target',$2) ON CONFLICT DO NOTHING",
                    campaign_id, str(target["chat_id"]),
                )
        due = await self.fetch(
            "SELECT * FROM segment_campaign_deliveries WHERE campaign_id=$1 AND status IN ('pending','failed') AND attempts<3 ORDER BY channel,recipient_key",
            campaign_id,
        )
        sent_users = sent_targets = failed = 0
        for row in due:
            msg, attempts, error = await self._send_text_retry(int(row["recipient_key"]), body)
            if msg:
                await self.execute(
                    "UPDATE segment_campaign_deliveries SET status='sent',attempts=attempts+$4,error=NULL,sent_at=NOW() WHERE campaign_id=$1 AND channel=$2 AND recipient_key=$3",
                    campaign_id, row["channel"], row["recipient_key"], attempts,
                )
                if row["channel"] == "dm":
                    sent_users += 1
                else:
                    sent_targets += 1
            else:
                failed += 1
                await self.execute(
                    "UPDATE segment_campaign_deliveries SET status='failed',attempts=attempts+$4,error=$5 WHERE campaign_id=$1 AND channel=$2 AND recipient_key=$3",
                    campaign_id, row["channel"], row["recipient_key"], attempts, str(error)[:500] if error else "unknown",
                )
            await asyncio.sleep(0.05)
        await self.execute(
            """UPDATE segment_campaigns SET status='sent',sent_at=COALESCE(sent_at,NOW()),
            sent_users=(SELECT COUNT(*) FROM segment_campaign_deliveries WHERE campaign_id=$1 AND channel='dm' AND status='sent'),
            sent_targets=(SELECT COUNT(*) FROM segment_campaign_deliveries WHERE campaign_id=$1 AND channel='target' AND status='sent'),
            failed=(SELECT COUNT(*) FROM segment_campaign_deliveries WHERE campaign_id=$1 AND status='failed') WHERE id=$1""",
            campaign_id,
        )

    async def _survey_due(self) -> None:
        now_local = datetime.now(SURVEY_TZ)
        if now_local.day < SURVEY_DAY or (now_local.day == SURVEY_DAY and now_local.hour < SURVEY_HOUR):
            return
        month_key = now_local.strftime("%Y-%m")
        existing = await self.fetchrow("SELECT * FROM survey_runs WHERE month_key=$1", month_key)
        if not existing:
            template = survey_template_for_month(now_local.year, now_local.month)
            scheduled = datetime(now_local.year, now_local.month, SURVEY_DAY, SURVEY_HOUR, tzinfo=SURVEY_TZ).astimezone(timezone.utc)
            run = await self.fetchrow(
                "INSERT INTO survey_runs(month_key,template_key,status,scheduled_for) VALUES($1,$2,'sending',$3) ON CONFLICT(month_key) DO NOTHING RETURNING *",
                month_key, template["key"], scheduled,
            )
            if run:
                audience = await self._active_audience_rows()
                for user in audience:
                    await self.execute(
                        "INSERT INTO survey_recipients(run_id,telegram_id) VALUES($1,$2) ON CONFLICT DO NOTHING",
                        run["id"], user["telegram_id"],
                    )
                await self.execute(
                    "UPDATE survey_runs SET target_count=(SELECT COUNT(*) FROM survey_recipients WHERE run_id=$1) WHERE id=$1",
                    run["id"],
                )
                existing = run
        if existing and existing["status"] in {"sending", "open"}:
            await self._deliver_survey_invites(existing["id"])

    async def _deliver_survey_invites(self, run_id: int) -> None:
        run = await self.fetchrow("SELECT * FROM survey_runs WHERE id=$1", run_id)
        if not run:
            return
        template = SURVEY_BY_KEY.get(run["template_key"])
        if not template:
            return
        due = await self.fetch(
            "SELECT * FROM survey_recipients WHERE run_id=$1 AND status IN ('pending','failed') AND attempts<3 ORDER BY telegram_id LIMIT 100",
            run_id,
        )
        for recipient in due:
            try:
                await self.bot.send_message(
                    recipient["telegram_id"],
                    f"<b>💬 2 минуты, которые реально влияют на движение</b>\n\n{html.escape(template['intro'])}\n\nЭто короткий ежемесячный пульс: 4 оценки и, если захочешь, один комментарий. Ответы нужны не для отчёта ради отчёта — они помогают команде понимать, что усиливать, а что менять.",
                    reply_markup=_kb([[("🚀 Начать", f"sv:start:{run_id}")], [("Не хочу ежемесячные опросы", "sv:optout")]]),
                )
                await self.execute(
                    "UPDATE survey_recipients SET status='sent',attempts=attempts+1,error=NULL,invited_at=NOW() WHERE run_id=$1 AND telegram_id=$2",
                    run_id, recipient["telegram_id"],
                )
            except Exception as exc:
                await self.execute(
                    "UPDATE survey_recipients SET status='failed',attempts=attempts+1,error=$3 WHERE run_id=$1 AND telegram_id=$2",
                    run_id, recipient["telegram_id"], str(exc)[:500],
                )
            await asyncio.sleep(0.05)
        await self.execute(
            """UPDATE survey_runs SET sent_count=(SELECT COUNT(*) FROM survey_recipients WHERE run_id=$1 AND status IN ('sent','completed')),
            completed_count=(SELECT COUNT(*) FROM survey_recipients WHERE run_id=$1 AND status='completed'),
            status=CASE WHEN EXISTS(SELECT 1 FROM survey_recipients WHERE run_id=$1 AND status IN ('pending','failed') AND attempts<3) THEN 'sending' ELSE 'open' END
            WHERE id=$1""",
            run_id,
        )

    async def _survey_question(self, chat_id: int, run_id: int, uid: int, index: int) -> None:
        run = await self.fetchrow("SELECT * FROM survey_runs WHERE id=$1", run_id)
        if not run:
            return
        template = SURVEY_BY_KEY.get(run["template_key"])
        if not template or index >= len(template["questions"]):
            return await self._survey_comment_prompt(chat_id, run_id, uid)
        question = template["questions"][index]
        rows = [[(str(score), f"sv:a:{run_id}:{index}:{score}") for score in range(1, 6)]]
        rows.append([("✖️ Выйти — продолжу позже", "menu")])
        await self.bot.send_message(
            chat_id,
            f"<b>{html.escape(template['title'])}</b> · {index + 1}/{len(template['questions'])}\n\n{html.escape(question)}\n\n1 — совсем не согласен(на) · 5 — полностью согласен(на)",
            reply_markup=_kb(rows),
        )

    async def _survey_comment_prompt(self, chat_id: int, run_id: int, uid: int) -> None:
        run = await self.fetchrow("SELECT * FROM survey_runs WHERE id=$1", run_id)
        template = SURVEY_BY_KEY.get(run["template_key"]) if run else None
        if not template:
            return
        await self.state_set(uid, "survey:comment", {"run_id": run_id})
        await self.bot.send_message(
            chat_id,
            f"<b>Последний штрих ✨</b>\n\n{html.escape(template['comment'])}\n\nМожно ответить одной фразой — или пропустить.",
            reply_markup=_kb([[("Пропустить →", f"sv:comment:skip:{run_id}")]]),
        )

    async def _complete_survey(self, chat_id: int, uid: int, run_id: int, comment: str | None) -> None:
        progress = await self.fetchrow("SELECT * FROM survey_progress WHERE run_id=$1 AND telegram_id=$2", run_id, uid)
        answers = _dict(progress["answers"]) if progress else {}
        if comment:
            answers["comment"] = comment[:1500]
        await self.execute(
            "INSERT INTO survey_responses(run_id,telegram_id,answers) VALUES($1,$2,$3::jsonb) ON CONFLICT(run_id,telegram_id) DO UPDATE SET answers=EXCLUDED.answers,completed_at=NOW()",
            run_id, uid, json.dumps(answers, ensure_ascii=False),
        )
        await self.execute("DELETE FROM survey_progress WHERE run_id=$1 AND telegram_id=$2", run_id, uid)
        await self.execute("UPDATE survey_recipients SET status='completed',completed_at=NOW() WHERE run_id=$1 AND telegram_id=$2", run_id, uid)
        await self.execute(
            "UPDATE survey_runs SET completed_count=(SELECT COUNT(*) FROM survey_recipients WHERE run_id=$1 AND status='completed') WHERE id=$1",
            run_id,
        )
        await self.state_clear(uid)
        await self.bot.send_message(
            chat_id,
            "<b>Спасибо. Именно так движение становится сильнее 💛</b>\n\nТвой ответ сохранён. В ежемесячной аналитике команда увидит общую картину и сможет принимать решения не на ощущениях, а на реальной обратной связи.",
        )

    async def _show_effectiveness(self, chat_id: int, user) -> None:
        if user["role"] not in {"admin", "owner"}:
            return
        latest = await self.fetchrow("SELECT * FROM survey_runs ORDER BY month_key DESC LIMIT 1")
        if not latest:
            next_template = survey_template_for_month(datetime.now(SURVEY_TZ).year, datetime.now(SURVEY_TZ).month)
            text = (
                "<b>📈 Эффективность движения</b>\n\n"
                "Ежемесячный пульс ещё не запускался. Бот автоматически отправляет один короткий опрос активной аудитории 5-го числа каждого месяца в 19:00 по Еревану.\n\n"
                f"Ближайшая тема: <b>{html.escape(next_template['title'])}</b>."
            )
        else:
            template = SURVEY_BY_KEY.get(latest["template_key"], {})
            sent = int(latest["sent_count"] or 0)
            completed = int(latest["completed_count"] or 0)
            rate = round(completed * 100 / sent, 1) if sent else 0
            text = (
                f"<b>📈 Эффективность движения · {latest['month_key']}</b>\n\n"
                f"Тема: <b>{html.escape(template.get('title', latest['template_key']))}</b>\n"
                f"🎯 Аудитория: {latest['target_count']}\n📨 Доставлено: {sent}\n✅ Ответили: {completed}\n📊 Отклик: <b>{rate}%</b>"
            )
        rows = [[("📊 Результаты", "svadm:results"), ("📥 Excel", "svadm:excel")], [("🧭 10 тем опросов", "svadm:templates")], [("⬅️ Управление", "menu")]]
        await self.bot.send_message(chat_id, text, reply_markup=_kb(rows))

    async def _survey_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Monthly pulse"
        ws.append(["Месяц", "Тема", "Аудитория", "Доставлено", "Ответили", "Отклик %", "Средняя оценка"])
        runs = await self.fetch("SELECT * FROM survey_runs ORDER BY month_key DESC")
        for run in runs:
            responses = await self.fetch("SELECT answers FROM survey_responses WHERE run_id=$1", run["id"])
            scores: list[int] = []
            for response in responses:
                data = _dict(response["answers"])
                scores.extend(int(v) for k, v in data.items() if k.startswith("q") and str(v).isdigit())
            avg = round(sum(scores) / len(scores), 2) if scores else None
            sent = int(run["sent_count"] or 0)
            completed = int(run["completed_count"] or 0)
            template = SURVEY_BY_KEY.get(run["template_key"], {})
            ws.append([run["month_key"], template.get("title", run["template_key"]), run["target_count"], sent, completed, round(completed * 100 / sent, 1) if sent else 0, avg])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
        details = wb.create_sheet("Responses")
        details.append(["Месяц", "Telegram ID", "Страна", "Возраст", "Q1", "Q2", "Q3", "Q4", "Комментарий", "Дата"])
        rows = await self.fetch(
            """SELECT sr.month_key,r.telegram_id,u.country_name,u.age,r.answers,r.completed_at
            FROM survey_responses r JOIN survey_runs sr ON sr.id=r.run_id JOIN users u ON u.telegram_id=r.telegram_id
            ORDER BY sr.month_key DESC,r.completed_at"""
        )
        for row in rows:
            answers = _dict(row["answers"])
            details.append([row["month_key"], row["telegram_id"], row["country_name"], row["age"], answers.get("q0"), answers.get("q1"), answers.get("q2"), answers.get("q3"), answers.get("comment"), row["completed_at"].replace(tzinfo=None) if row["completed_at"] else None])
        details.freeze_panes = "A2"
        details.auto_filter.ref = details.dimensions
        for cell in details[1]:
            cell.font = Font(bold=True)
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def _engagement_scheduler(self) -> None:
        while True:
            try:
                await self._survey_due()
                failed = await self.fetch("SELECT DISTINCT campaign_id FROM segment_campaign_deliveries WHERE status='failed' AND attempts<3 LIMIT 10")
                for row in failed:
                    await self._deliver_segment_campaign(row["campaign_id"])
            except Exception:
                log.exception("engagement scheduler")
            await asyncio.sleep(60)

    async def reminders_loop(self):
        task = asyncio.create_task(self._engagement_scheduler())
        try:
            await super().reminders_loop()
        finally:
            task.cancel()

    def _register_handlers(self):
        r = self.router

        @r.callback_query(F.data == "seg:menu")
        async def segment_menu(c: CallbackQuery):
            await c.answer()
            user = await self.ensure_user(c.from_user)
            if user["role"] not in {"admin", "owner"}:
                return
            recent = await self.fetch("SELECT * FROM segment_campaigns ORDER BY id DESC LIMIT 5")
            lines = ["<b>🎯 Рассылки и сегменты</b>", "", "Отправляй точные сообщения нужным людям и нужным чатам. Перед отправкой бот всегда показывает размер аудитории и просит финальное подтверждение."]
            if recent:
                lines.append("\n<b>Последние:</b>")
                for item in recent:
                    lines.append(f"#{item['id']} · {item['status']} · 👤 {item['sent_users']} · 📣 {item['sent_targets']} · ⚠️ {item['failed']}")
            await c.message.answer("\n".join(lines), reply_markup=_kb([[("➕ Новая рассылка", "seg:new")], [("📣 Настроить чаты", "tg:list")], [("⬅️ Управление", "menu")]]))

        @r.callback_query(F.data == "seg:new")
        async def segment_new(c: CallbackQuery):
            await c.answer()
            user = await self.ensure_user(c.from_user)
            if user["role"] not in {"admin", "owner"}:
                return
            p = {"segment": {}, "channels": {"dm": True, "targets": False}}
            await self._show_segment_builder(c.message.chat.id, user["telegram_id"], p)

        @r.callback_query(F.data == "seg:geo")
        async def segment_geo(c: CallbackQuery):
            await c.answer()
            await c.message.answer("<b>География</b>\n\nМожно выбрать весь мир, один или несколько регионов либо конкретные страны.", reply_markup=_kb([[("🌍 Весь мир", "seg:g:all")], [("🗺 Регионы", "seg:g:regions"), ("🏳 Страны", "seg:g:countries")], [("⬅️ Назад", "seg:back")]]))

        @r.callback_query(F.data == "seg:g:all")
        async def segment_geo_all(c: CallbackQuery):
            await c.answer()
            st = await self.state_get(c.from_user.id)
            p = dict(st["payload"] or {}) if st else {"segment": {}, "channels": {"dm": True}}
            p.setdefault("segment", {})["countries"] = []
            p["segment"]["regions"] = []
            await self._show_segment_builder(c.message.chat.id, c.from_user.id, p)

        @r.callback_query(F.data == "seg:g:regions")
        async def segment_regions(c: CallbackQuery):
            await c.answer()
            rows = []
            for i in range(0, len(REGION_LABELS), 2):
                pairs = list(REGION_LABELS.items())[i:i + 2]
                rows.append([(label, f"seg:r:{key}") for key, label in pairs])
            rows.append([("Готово →", "seg:back")])
            await c.message.answer("Выбери один или несколько регионов:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("seg:r:"))
        async def segment_region_toggle(c: CallbackQuery):
            await c.answer()
            key = c.data.split(":", 2)[2]
            if key not in REGION_COUNTRIES:
                return
            st = await self.state_get(c.from_user.id)
            p = dict(st["payload"] or {})
            seg = p.setdefault("segment", {})
            selected = list(seg.get("regions") or [])
            selected = [x for x in selected if x != key] if key in selected else selected + [key]
            seg["regions"] = selected
            seg["countries"] = []
            await self.state_set(c.from_user.id, "segment:builder", p)
            await c.message.answer(f"{'✅ Добавлен' if key in selected else 'Убран'} регион: {REGION_LABELS[key]}")

        @r.callback_query(F.data == "seg:g:countries")
        async def segment_country_regions(c: CallbackQuery):
            await c.answer()
            rows = []
            items = list(REGION_LABELS.items())
            for i in range(0, len(items), 2):
                rows.append([(label, f"seg:cr:{key}") for key, label in items[i:i + 2]])
            rows.append([("⬅️ Назад", "seg:back")])
            await c.message.answer("Сначала выбери регион, чтобы быстро найти страны:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("seg:cr:"))
        async def segment_country_region(c: CallbackQuery):
            await c.answer()
            region = c.data.split(":", 2)[2]
            if region not in REGION_COUNTRIES:
                return
            codes = REGION_COUNTRIES[region]
            rows = []
            for i in range(0, len(codes), 2):
                rows.append([(_country_label(code), f"seg:c:{code}") for code in codes[i:i + 2]])
            rows.append([("Готово →", "seg:back")])
            await c.message.answer(f"<b>{REGION_LABELS[region]}</b>\nМожно отметить несколько стран:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("seg:c:"))
        async def segment_country_toggle(c: CallbackQuery):
            await c.answer()
            code = c.data.split(":", 2)[2]
            if code not in WORLD_COUNTRIES:
                return
            st = await self.state_get(c.from_user.id)
            p = dict(st["payload"] or {})
            seg = p.setdefault("segment", {})
            selected = list(seg.get("countries") or [])
            selected = [x for x in selected if x != code] if code in selected else selected + [code]
            seg["countries"] = selected
            seg["regions"] = []
            await self.state_set(c.from_user.id, "segment:builder", p)
            await c.answer("Добавлено" if code in selected else "Убрано", show_alert=False)

        @r.callback_query(F.data == "seg:age")
        async def segment_age(c: CallbackQuery):
            await c.answer()
            await c.message.answer("Возрастные группы — можно выбрать несколько:", reply_markup=_kb([[((label), f"seg:a:{key}") for key, (_, _, label) in AGE_BUCKETS.items()], [("Готово →", "seg:back")]]))

        @r.callback_query(F.data.startswith("seg:a:"))
        async def segment_age_toggle(c: CallbackQuery):
            await c.answer()
            key = c.data.split(":", 2)[2]
            if key not in AGE_BUCKETS:
                return
            st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}); seg = p.setdefault("segment", {})
            selected = list(seg.get("age_buckets") or []); seg["age_buckets"] = [x for x in selected if x != key] if key in selected else selected + [key]
            await self.state_set(c.from_user.id, "segment:builder", p)

        @r.callback_query(F.data == "seg:status")
        async def segment_status(c: CallbackQuery):
            await c.answer()
            items = list(STATUS_LABELS.items()); rows = []
            for i in range(0, len(items), 2): rows.append([(label, f"seg:s:{key}") for key, label in items[i:i + 2]])
            rows.append([("Готово →", "seg:back")]); await c.message.answer("Статус участника:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("seg:s:"))
        async def segment_status_toggle(c: CallbackQuery):
            await c.answer(); key = c.data.split(":", 2)[2]
            if key not in STATUS_LABELS: return
            st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}); seg = p.setdefault("segment", {})
            selected = list(seg.get("statuses") or []); seg["statuses"] = [x for x in selected if x != key] if key in selected else selected + [key]
            await self.state_set(c.from_user.id, "segment:builder", p)

        @r.callback_query(F.data == "seg:interest")
        async def segment_interest(c: CallbackQuery):
            await c.answer(); items = list(INTEREST_LABELS.items()); rows = []
            for i in range(0, len(items), 2): rows.append([(label, f"seg:i:{key}") for key, label in items[i:i + 2]])
            rows.append([("Готово →", "seg:back")]); await c.message.answer("Интересы аудитории:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("seg:i:"))
        async def segment_interest_toggle(c: CallbackQuery):
            await c.answer(); key = c.data.split(":", 2)[2]
            if key not in INTEREST_LABELS: return
            st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}); seg = p.setdefault("segment", {})
            selected = list(seg.get("interests") or []); seg["interests"] = [x for x in selected if x != key] if key in selected else selected + [key]
            await self.state_set(c.from_user.id, "segment:builder", p)

        @r.callback_query(F.data == "seg:active")
        async def segment_active(c: CallbackQuery):
            await c.answer(); st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}); seg = p.setdefault("segment", {}); seg["active_only"] = not bool(seg.get("active_only")); await self._show_segment_builder(c.message.chat.id, c.from_user.id, p)

        @r.callback_query(F.data.startswith("seg:ch:"))
        async def segment_channel(c: CallbackQuery):
            await c.answer(); channel = c.data.split(":", 2)[2]
            st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}); channels = p.setdefault("channels", {"dm": True, "targets": False}); channels[channel] = not bool(channels.get(channel)); await self._show_segment_builder(c.message.chat.id, c.from_user.id, p)

        @r.callback_query(F.data.in_({"seg:back", "seg:estimate"}))
        async def segment_back(c: CallbackQuery):
            await c.answer(); st = await self.state_get(c.from_user.id); p = dict(st["payload"] or {}) if st else {"segment": {}, "channels": {"dm": True}}; await self._show_segment_builder(c.message.chat.id, c.from_user.id, p)

        @r.callback_query(F.data == "seg:compose")
        async def segment_compose(c: CallbackQuery):
            await c.answer(); st = await self.state_get(c.from_user.id)
            if not st: return
            p = dict(st["payload"] or {}); users_count, targets_count = await self._estimate_segment(p.get("segment", {}), p.get("channels", {}))
            if users_count + targets_count == 0: return await c.message.answer("По этим условиям аудитория пустая. Измени сегмент — отправка не запустится в пустоту.")
            p["estimate"] = {"users": users_count, "targets": targets_count}; await self.state_set(c.from_user.id, "segment:text", p)
            await c.message.answer("<b>Сообщение для аудитории</b>\n\nОтправь текст одним сообщением. До 3500 символов. После этого покажу финальный предпросмотр — ничего не уйдёт без подтверждения.")

        @r.callback_query(F.data == "seg:send")
        async def segment_send(c: CallbackQuery):
            await c.answer(); user = await self.ensure_user(c.from_user)
            if user["role"] not in {"admin", "owner"}: return
            st = await self.state_get(user["telegram_id"])
            if not st or st["state"] != "segment:preview": return
            p = dict(st["payload"] or {}); users_count, targets_count = await self._estimate_segment(p.get("segment", {}), p.get("channels", {}))
            if users_count + targets_count == 0: return await c.message.answer("Аудитория стала пустой. Рассылка остановлена.")
            campaign = await self.fetchrow(
                """INSERT INTO segment_campaigns(created_by,body,segment,channels,status,estimated_users,estimated_targets)
                VALUES($1,$2,$3::jsonb,$4::jsonb,'sending',$5,$6) RETURNING *""",
                user["telegram_id"], p["body"], json.dumps(p.get("segment", {}), ensure_ascii=False), json.dumps(p.get("channels", {}), ensure_ascii=False), users_count, targets_count,
            )
            await self.state_clear(user["telegram_id"]); await self.audit(user["telegram_id"], "segment_campaign_started", "segment_campaign", campaign["id"], {"users": users_count, "targets": targets_count})
            await c.message.answer(f"Рассылка #{campaign['id']} запущена ✅\nАудитория: {users_count} участников + {targets_count} чатов/каналов.")
            asyncio.create_task(self._deliver_segment_campaign(campaign["id"]))

        @r.callback_query(F.data == "seg:cancel")
        async def segment_cancel(c: CallbackQuery):
            await c.answer(); await self.state_clear(c.from_user.id); await c.message.answer("Рассылка отменена. Ничего не отправлено.")

        @r.callback_query(F.data == "tg:list")
        async def target_list(c: CallbackQuery):
            await c.answer(); user = await self.ensure_user(c.from_user)
            if user["role"] not in {"admin", "owner"}: return
            rows = await self.fetch("SELECT * FROM targets ORDER BY CASE status WHEN 'pending' THEN 0 WHEN 'approved' THEN 1 ELSE 2 END,updated_at DESC LIMIT 50")
            buttons = [[(("🟡" if t["status"] == "pending" else "🟢" if t["status"] == "approved" else "⚫") + " " + t["title"][:48], f"tg:open:{t['chat_id']}")] for t in rows]
            buttons += [[("➕ Как подключить", "tg:guide")], [("⬅️ Управление", "menu")]]
            await c.message.answer("<b>📣 Чаты и каналы</b>\n\nКаждый канал получает понятную географию, назначение и теги. Благодаря этому рассылка не попадёт случайно не той аудитории.", reply_markup=_kb(buttons))

        @r.callback_query(F.data == "tg:guide")
        async def target_guide(c: CallbackQuery):
            await c.answer(); await c.message.answer("<b>Как подключить чат или канал</b>\n\n1. Добавь @MVKSRS_bot в группу или канал.\n2. Выдай права администратора. Для канала обязательно разреши публикацию сообщений.\n3. Вернись сюда → 📣 Чаты и каналы.\n4. Новый объект будет отмечен 🟡.\n5. Укажи сеть, географию, назначение и при необходимости теги.\n6. Нажми «Подтвердить».\n\nПосле этого чат становится доступен для точной сегментированной рассылки.")

        @r.callback_query(F.data.startswith("tg:open:"))
        async def target_open(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); target = await self.fetchrow("SELECT * FROM targets WHERE chat_id=$1", tid)
            if not target: return
            tags = _json_list(target["tags"]); geo = "🌍 весь мир" if target["geography_scope"] == "global" else REGION_LABELS.get(target["region_code"], "") if target["geography_scope"] == "region" else target["country_code"] or "не настроена"
            text = f"<b>{html.escape(target['title'])}</b>\n\nСтатус: {target['status']}\nСеть: {target['network']}\nГеография: {geo}\nНазначение: {TARGET_PURPOSES.get(target['purpose'], target['purpose'])}\nТеги: {', '.join(TARGET_TAGS.get(x,x) for x in tags) or '—'}\nПрава публикации: {'✅' if target['bot_can_post'] else '❌'}"
            await c.message.answer(text, reply_markup=_kb([[("🌐 Сеть", f"tg:nets:{tid}"), ("🌍 География", f"tg:geo:{tid}")], [("🧭 Назначение", f"tg:purpose:{tid}"), ("🏷 Теги", f"tg:tags:{tid}")], [("✅ Подтвердить", f"tg:approve:{tid}")], [("⬅️ К списку", "tg:list")]]))

        @r.callback_query(F.data.startswith("tg:nets:"))
        async def target_nets(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); await c.message.answer("К какой сети относится площадка?", reply_markup=_kb([[('Комиссия', f"tg:net:{tid}:commission"), ('МДС', f"tg:net:{tid}:mds")], [('Партнёр', f"tg:net:{tid}:partner")]]))

        @r.callback_query(F.data.startswith("tg:net:"))
        async def target_net(c: CallbackQuery):
            await c.answer(); _, _, tid, network = c.data.split(":", 3); await self.execute("UPDATE targets SET network=$2,updated_at=NOW() WHERE chat_id=$1", int(tid), network); await c.message.answer("Сеть сохранена ✅", reply_markup=_kb([[("Продолжить настройку", f"tg:open:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:geo:"))
        async def target_geo(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); await c.message.answer("География этой площадки:", reply_markup=_kb([[('🌍 Международная', f"tg:g:{tid}:global")], [('🗺 Регион', f"tg:gr:{tid}"), ('🏳 Страна', f"tg:gc:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:g:"))
        async def target_global(c: CallbackQuery):
            await c.answer(); _, _, tid, _ = c.data.split(":", 3); await self.execute("UPDATE targets SET geography_scope='global',region_code=NULL,country_code=NULL,updated_at=NOW() WHERE chat_id=$1", int(tid)); await c.message.answer("География: весь мир ✅", reply_markup=_kb([[("Продолжить", f"tg:open:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:gr:"))
        async def target_region_menu(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); items = list(REGION_LABELS.items()); rows=[]
            for i in range(0,len(items),2): rows.append([(label, f"tg:r:{tid}:{key}") for key,label in items[i:i+2]])
            await c.message.answer("Выбери регион:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("tg:r:"))
        async def target_region_set(c: CallbackQuery):
            await c.answer(); _, _, tid, region = c.data.split(":", 3)
            if region not in REGION_COUNTRIES: return
            await self.execute("UPDATE targets SET geography_scope='region',region_code=$2,country_code=NULL,updated_at=NOW() WHERE chat_id=$1", int(tid), region); await c.message.answer("Регион сохранён ✅", reply_markup=_kb([[("Продолжить", f"tg:open:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:gc:"))
        async def target_country_region_menu(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); items = list(REGION_LABELS.items()); rows=[]
            for i in range(0,len(items),2): rows.append([(label, f"tg:cr:{tid}:{key}") for key,label in items[i:i+2]])
            await c.message.answer("Сначала регион:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("tg:cr:"))
        async def target_country_menu(c: CallbackQuery):
            await c.answer(); _, _, tid, region = c.data.split(":", 3)
            if region not in REGION_COUNTRIES: return
            codes = REGION_COUNTRIES[region]; rows=[]
            for i in range(0,len(codes),2): rows.append([(_country_label(code), f"tg:c:{tid}:{code}") for code in codes[i:i+2]])
            await c.message.answer(f"<b>{REGION_LABELS[region]}</b>\nВыбери страну:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("tg:c:"))
        async def target_country_set(c: CallbackQuery):
            await c.answer(); _, _, tid, code = c.data.split(":", 3)
            if code not in WORLD_COUNTRIES: return
            await self.execute("UPDATE targets SET geography_scope='country',country_code=$2,region_code=$3,updated_at=NOW() WHERE chat_id=$1", int(tid), code, country_region(code)); await c.message.answer("Страна сохранена ✅", reply_markup=_kb([[("Продолжить", f"tg:open:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:purpose:"))
        async def target_purpose(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); items=list(TARGET_PURPOSES.items()); rows=[]
            for i in range(0,len(items),2): rows.append([(label, f"tg:p:{tid}:{key}") for key,label in items[i:i+2]])
            await c.message.answer("Для каких публикаций использовать эту площадку?", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("tg:p:"))
        async def target_purpose_set(c: CallbackQuery):
            await c.answer(); _, _, tid, purpose = c.data.split(":", 3)
            if purpose not in TARGET_PURPOSES: return
            await self.execute("UPDATE targets SET purpose=$2,updated_at=NOW() WHERE chat_id=$1", int(tid), purpose); await c.message.answer("Назначение сохранено ✅", reply_markup=_kb([[("Продолжить", f"tg:open:{tid}")]]))

        @r.callback_query(F.data.startswith("tg:tags:"))
        async def target_tags(c: CallbackQuery):
            await c.answer(); tid = int(c.data.split(":", 2)[2]); target=await self.fetchrow("SELECT tags FROM targets WHERE chat_id=$1", tid); selected=set(_json_list(target["tags"]) if target else []); items=list(TARGET_TAGS.items()); rows=[]
            for i in range(0,len(items),2): rows.append([(("✅ " if key in selected else "▫️ ")+label, f"tg:t:{tid}:{key}") for key,label in items[i:i+2]])
            rows.append([("Готово →", f"tg:open:{tid}")]); await c.message.answer("Теги помогают ещё точнее выбирать площадки:", reply_markup=_kb(rows))

        @r.callback_query(F.data.startswith("tg:t:"))
        async def target_tag_toggle(c: CallbackQuery):
            await c.answer(); _, _, tid, tag = c.data.split(":", 3)
            if tag not in TARGET_TAGS: return
            target=await self.fetchrow("SELECT tags FROM targets WHERE chat_id=$1", int(tid)); selected=_json_list(target["tags"]) if target else []; selected=[x for x in selected if x!=tag] if tag in selected else selected+[tag]
            await self.execute("UPDATE targets SET tags=$2::jsonb,updated_at=NOW() WHERE chat_id=$1", int(tid), json.dumps(selected)); await c.answer("Сохранено")

        @r.callback_query(F.data.startswith("tg:approve:"))
        async def target_approve(c: CallbackQuery):
            await c.answer(); user=await self.ensure_user(c.from_user)
            if user["role"] not in {"admin","owner"}: return
            tid=int(c.data.split(":",2)[2]); target=await self.fetchrow("SELECT * FROM targets WHERE chat_id=$1",tid)
            if not target or not target["bot_can_post"]: return await c.message.answer("Не могу подтвердить: у бота нет прав на публикацию.")
            if target["geography_scope"] == "country" and not target["country_code"]: return await c.message.answer("Сначала укажи страну или выбери регион/международную географию.")
            if target["geography_scope"] == "region" and not target["region_code"]: return await c.message.answer("Сначала укажи регион.")
            await self.execute("UPDATE targets SET status='approved',approved_by=$2,approved_at=NOW(),updated_at=NOW() WHERE chat_id=$1",tid,user["telegram_id"]); await self.audit(user["telegram_id"],"target_approved","target",tid); await c.message.answer("Площадка подключена ✅ Теперь она участвует только в подходящих сегментах.")

        @r.callback_query(F.data == "svadm:menu")
        async def survey_admin_menu(c: CallbackQuery):
            await c.answer(); user=await self.ensure_user(c.from_user); await self._show_effectiveness(c.message.chat.id,user)

        @r.callback_query(F.data == "svadm:templates")
        async def survey_templates(c: CallbackQuery):
            await c.answer(); lines=["<b>10 ежемесячных пульсов</b>",""]
            for idx,item in enumerate(SURVEY_TEMPLATES,1): lines.append(f"{idx}. <b>{html.escape(item['title'])}</b> — {html.escape(item['intro'])}")
            await c.message.answer("\n\n".join(lines))

        @r.callback_query(F.data == "svadm:results")
        async def survey_results(c: CallbackQuery):
            await c.answer(); run=await self.fetchrow("SELECT * FROM survey_runs ORDER BY month_key DESC LIMIT 1")
            if not run: return await c.message.answer("Ответов пока нет.")
            template=SURVEY_BY_KEY.get(run["template_key"]); responses=await self.fetch("SELECT answers FROM survey_responses WHERE run_id=$1",run["id"]); lines=[f"<b>{html.escape(template['title'])}</b> · {run['month_key']}"]
            for i,q in enumerate(template["questions"]):
                vals=[]
                for rr in responses:
                    v=_dict(rr["answers"]).get(f"q{i}")
                    if isinstance(v,int) or str(v).isdigit(): vals.append(int(v))
                lines.append(f"\n{i+1}. {html.escape(q)}\n<b>{round(sum(vals)/len(vals),2) if vals else '—'}</b> / 5 · ответов {len(vals)}")
            await c.message.answer("\n".join(lines),reply_markup=_kb([[("📥 Excel", "svadm:excel")]]))

        @r.callback_query(F.data == "svadm:excel")
        async def survey_excel(c: CallbackQuery):
            await c.answer("Готовлю Excel…"); user=await self.ensure_user(c.from_user)
            if user["role"] not in {"admin","owner"}: return
            data=await self._survey_xlsx(); stamp=datetime.now().strftime("%Y-%m-%d_%H-%M"); await self.bot.send_document(c.message.chat.id,BufferedInputFile(data,filename=f"movement_effectiveness_{stamp}.xlsx"),caption="📈 Ежемесячная эффективность движения: динамика и ответы.")

        @r.callback_query(F.data.startswith("sv:start:"))
        async def survey_start(c: CallbackQuery):
            await c.answer(); run_id=int(c.data.split(":",2)[2]); user=await self.ensure_user(c.from_user)
            recipient=await self.fetchrow("SELECT * FROM survey_recipients WHERE run_id=$1 AND telegram_id=$2",run_id,user["telegram_id"])
            if not recipient: return await c.message.answer("Этот опрос не назначен вашему профилю.")
            done=await self.fetchrow("SELECT 1 FROM survey_responses WHERE run_id=$1 AND telegram_id=$2",run_id,user["telegram_id"])
            if done: return await c.message.answer("Ты уже ответил(а) на этот пульс. Спасибо 💛")
            progress=await self.fetchrow("SELECT * FROM survey_progress WHERE run_id=$1 AND telegram_id=$2",run_id,user["telegram_id"])
            index=int(progress["question_index"] if progress else 0); await self._survey_question(c.message.chat.id,run_id,user["telegram_id"],index)

        @r.callback_query(F.data.startswith("sv:a:"))
        async def survey_answer(c: CallbackQuery):
            await c.answer(); _,_,run_s,index_s,score_s=c.data.split(":",4); run_id=int(run_s); index=int(index_s); score=int(score_s)
            if score not in range(1,6): return
            run=await self.fetchrow("SELECT * FROM survey_runs WHERE id=$1",run_id); template=SURVEY_BY_KEY.get(run["template_key"]) if run else None
            if not template or index>=len(template["questions"]): return
            progress=await self.fetchrow("SELECT * FROM survey_progress WHERE run_id=$1 AND telegram_id=$2",run_id,c.from_user.id); answers=_dict(progress["answers"]) if progress else {}; answers[f"q{index}"]=score
            await self.execute("""INSERT INTO survey_progress(run_id,telegram_id,question_index,answers) VALUES($1,$2,$3,$4::jsonb)
                ON CONFLICT(run_id,telegram_id) DO UPDATE SET question_index=EXCLUDED.question_index,answers=EXCLUDED.answers,updated_at=NOW()""",run_id,c.from_user.id,index+1,json.dumps(answers,ensure_ascii=False))
            if index+1>=len(template["questions"]): await self._survey_comment_prompt(c.message.chat.id,run_id,c.from_user.id)
            else: await self._survey_question(c.message.chat.id,run_id,c.from_user.id,index+1)

        @r.callback_query(F.data.startswith("sv:comment:skip:"))
        async def survey_skip_comment(c: CallbackQuery):
            await c.answer(); run_id=int(c.data.rsplit(":",1)[1]); await self._complete_survey(c.message.chat.id,c.from_user.id,run_id,None)

        @r.callback_query(F.data == "sv:optout")
        async def survey_optout(c: CallbackQuery):
            await c.answer(); await self.execute("UPDATE users SET monthly_survey_opt_in=FALSE WHERE telegram_id=$1",c.from_user.id); await c.message.answer("Ежемесячные опросы выключены. Обычные уведомления и регистрации продолжат работать как раньше.")

        @r.message(F.chat.type == ChatType.PRIVATE)
        async def engagement_text_state(m: Message):
            user=await self.ensure_user(m.from_user); st=await self.state_get(user["telegram_id"])
            if not st: raise SkipHandler
            state=str(st["state"]); p=dict(st["payload"] or {}); text=(m.text or "").strip()
            if state == "segment:text":
                if not text: return await m.answer("Отправь текст сообщением.")
                if len(text)>3500: return await m.answer("Сделай сообщение короче 3500 символов — так оно надёжно отправится во все Telegram-площадки.")
                p["body"]=text; users_count,targets_count=await self._estimate_segment(p.get("segment",{}),p.get("channels",{})); p["estimate"]={"users":users_count,"targets":targets_count}; await self.state_set(user["telegram_id"],"segment:preview",p)
                return await m.answer(f"<b>Финальная проверка</b>\n\n{html.escape(text)}\n\n{self._segment_summary(p)}\n\n👥 Участники: <b>{users_count}</b>\n📣 Чаты/каналы: <b>{targets_count}</b>\n\nПосле кнопки «Отправить» рассылка начнётся сразу.",reply_markup=_kb([[("🚀 Отправить", "seg:send")],[("✖️ Отмена", "seg:cancel")]]))
            if state == "survey:comment":
                run_id=int(p.get("run_id") or 0)
                if not run_id: await self.state_clear(user["telegram_id"]); raise SkipHandler
                return await self._complete_survey(m.chat.id,user["telegram_id"],run_id,text or None)
            raise SkipHandler

        super()._register_handlers()
