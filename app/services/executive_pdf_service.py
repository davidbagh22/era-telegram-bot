"""Organizational report using shared workbook facts, without personal lists."""
from datetime import date, datetime
from html import escape
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, CondPageBreak, KeepTogether, Table, TableStyle
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from app.services.career_pdf_service import _register_font

STATUS = {'draft': 'Черновик', 'initial_review': 'На рассмотрении', 'review': 'На проверке',
    'approved': 'Одобрен', 'in_progress': 'В работе', 'completed': 'Завершён',
    'cancelled': 'Отменён', 'rejected': 'Отклонён', 'pending': 'Ожидает',
    'published': 'Опубликовано', 'registration_open': 'Регистрация открыта',
    'registration_closed': 'Регистрация закрыта', 'report_submitted': 'Отчёт подан',
    'new': 'Новый', 'overdue': 'Просрочено', 'done': 'Выполнено', 'active': 'Активно'}


def build_executive_pdf(workbook: bytes, *, period_label: str, health, efficiency, analytics=None, media=None, telegram=None) -> bytes:
    _register_font()
    ink, muted, accent = (colors.HexColor(value) for value in ('#202126', '#71717B', '#632CFF'))
    styles = {key: ParagraphStyle(key, fontName='ERAUnicode', fontSize=size, leading=leading,
        textColor=color, spaceAfter=after, splitLongWords=True) for key, size, leading, color, after in [
        ('title', 30, 38, ink, 12), ('h', 18, 24, ink, 12), ('item', 12, 17, ink, 5),
        ('body', 10, 15, ink, 6), ('small', 8, 12, muted, 5), ('value', 26, 32, ink, 5),
        ('eyebrow', 9, 13, accent, 8)]}
    styles['h'].keepWithNext = True
    styles['item'].keepWithNext = True

    def text(value):
        if value is None:
            return 'Нет данных'
        return value.strftime('%d.%m.%Y') if isinstance(value, (date, datetime)) else str(value)

    def p(value, style='body'):
        return Paragraph(escape(text(value)).replace('\n', '<br/>'), styles[style])

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=20*mm, leftMargin=20*mm,
        topMargin=24*mm, bottomMargin=22*mm, title='ЭРА - Отчёт о деятельности', author='ЭРА')

    def page(canvas, document):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors.HexColor('#DADAE0'))
        canvas.line(20*mm, height-17*mm, width-20*mm, height-17*mm)
        canvas.setFont('ERAUnicode', 8)
        canvas.setFillColor(muted)
        canvas.drawString(20*mm, height-12*mm, 'ЭРА / ОТЧЁТ О ДЕЯТЕЛЬНОСТИ')
        canvas.drawString(20*mm, 12*mm, 'Для внутреннего использования · Без персональных списков')
        canvas.drawRightString(width-20*mm, 12*mm, f'{document.page:02d}')
        canvas.restoreState()

    wb = load_workbook(BytesIO(workbook), read_only=True, data_only=False)
    sheets = {sheet.title.split(' · ')[0]: sheet for sheet in wb.worksheets}

    def records(key):
        sheet = sheets.get(key)
        if sheet is None:
            return []
        values = iter(sheet.values)
        headers = next(values, ())
        return [dict(zip(headers, row)) for row in values if row and row[0] is not None
            and str(row[0]) != 'Данных за выбранный период нет']

    summary = {row[0]: (row[1], row[2]) for row in list(wb.worksheets[0].values)[4:] if len(row) >= 3 and row[0]}
    projects, events, tasks = records('08'), records('07'), records('09')
    story = [p('ИТОГИ ВЫБРАННОГО ПЕРИОДА', 'eyebrow'), p('Что сделала ЭРА', 'title'),
        p(period_label, 'h'), p('Проекты, мероприятия, задания и коммуникации. Только факты платформы; без анкет и списков людей.', 'small'), Spacer(1, 5*mm)]
    tiles = [[p(summary.get(name, ('Нет данных', 'Источник не предоставлен'))[0], 'value'), p(name),
        p(summary.get(name, ('', 'Источник не предоставлен'))[1], 'small')]
        for name in ['Текущий состав', 'Новые за период', 'События за период', 'Проекты за период']]
    table = Table([tiles[:2], tiles[2:]], colWidths=[85*mm, 85*mm], hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F4F3F7')),
        ('VALIGN',(0,0),(-1,-1),'TOP'), ('INNERGRID',(0,0),(-1,-1),5,colors.white),
        ('LEFTPADDING',(0,0),(-1,-1),12),('RIGHTPADDING',(0,0),(-1,-1),12),
        ('TOPPADDING',(0,0),(-1,-1),12),('BOTTOMPADDING',(0,0),(-1,-1),12)]))
    story.extend([table, Spacer(1, 6*mm), p('Результаты в одном абзаце', 'h'),
        p(f'В отчёт вошли: {len(projects)} проектов с созданием или подтверждённой работой, '
          f'{len(events)} мероприятий по календарной дате и {len(tasks)} новых заданий. '
          f'Подтверждённых посещений: {summary.get("Подтверждённых посещений", ("Нет данных",))[0]}. '
          f'Принятых результатов заданий: {summary.get("Verified task results", ("Нет данных",))[0]}.'),
        p('Ниже - перечни с датами и результатами. Статусы записей показаны на момент формирования, а не на конец выбранного месяца.', 'small'), PageBreak()])

    def section(title, note):
        story.extend([Spacer(1, 5*mm), CondPageBreak(45*mm), p(title, 'h'), p(note, 'small')])

    section('01 / Проекты за период', 'Новые проекты и проекты с подтверждённым вкладом или завершённым этапом за выбранные даты. Без авторов и списков участников.')
    if not projects:
        story.append(p('Проектов с зафиксированной работой за выбранные даты нет.'))
    for i, row in enumerate(projects, 1):
        status = row.get('Текущий статус', row.get('Статус'))
        story.extend([CondPageBreak(29*mm), p(f'{i:02d}. {row.get("Проект", "Проект")}', 'item'),
            p(f'{STATUS.get(status, text(status))} · создан {text(row.get("Создан"))}', 'small'),
            p(f'Подтверждённых вкладов: {row.get("Подтверждённых вкладов за период", 0)}. Завершено этапов: {row.get("Завершено этапов за период", 0)}.'),
            p(row.get('Основание включения', 'Создан в выбранном периоде'), 'small'), Spacer(1, 3*mm)])
    section('02 / Мероприятия', 'Перечень по дате мероприятия. Регистрация и фактическое посещение - разные показатели.')
    if not events:
        story.append(p('Мероприятий на выбранные даты в платформе нет.'))
    for i, row in enumerate(events, 1):
        status = row.get('Статус')
        story.extend([CondPageBreak(25*mm), p(f'{i:02d}. {row.get("Событие", "Мероприятие")}', 'item'),
            p(f'{text(row.get("Дата"))} · {STATUS.get(status, text(status))}', 'small'),
            p(f'Зарегистрировано: {text(row.get("Регистраций"))}. Посетили: {text(row.get("Посетили"))}.'), Spacer(1, 3*mm)])
    section('03 / Задания', 'Созданы за выбранные даты. Без исполнителей и текстов персональных ответов.')
    if not tasks:
        story.append(p('Новых заданий за выбранные даты нет.'))
    for i, row in enumerate(tasks, 1):
        status = row.get('Статус')
        story.extend([CondPageBreak(22*mm), p(f'{i:02d}. {row.get("Задание", "Задание")}', 'item'),
            p(f'{STATUS.get(status, text(status))} · создано {text(row.get("Создано"))} · срок {text(row.get("Дедлайн"))}', 'small')])

    telegram_start = len(story)
    section('04 / Канал и общий чат', 'Фактически полученные сообщения Telegram, а не план публикаций. Тексты и личности авторов не сохраняются в счётчике.')
    if not telegram or not telegram.get('sources'):
        story.append(p('Источники Telegram в этом отчёте не предоставлены.'))
    for index, source in enumerate((telegram or {}).get('sources', [])):
        source_start = telegram_start if index == 0 else len(story)
        def source_date(value):
            return datetime.fromisoformat(value).astimezone(ZoneInfo(telegram.get('timezone', 'Asia/Yerevan'))).strftime('%d.%m.%Y %H:%M') if value else 'ещё нет'
        story.extend([p(source['label'], 'item'), p(f'{source["observed_value"]} учтено', 'value'),
            p('Полный период' if source['data_state'] == 'valid' else 'Неполная история: учтённое количество, не полный итог периода.', 'small'),
            p(f'Подключение: {"сбор включён" if source.get("connection_status") == "collecting" else "требует проверки"}. Последнее поступление: {source_date(source.get("last_event_at"))}.', 'small'),
            p(f'Непрерывное покрытие с {source_date(source.get("coverage_since"))}.', 'small')])
        if source.get('rolling_counts'):
            story.append(p('Последние периоды на дату формирования: ' + ' / '.join(f'{item["days"]} дней: {item["observed"]} учтено' for item in source['rolling_counts'])))
        if source.get('delta') is not None:
            story.append(p(f'Изменение к предыдущему равному периоду: {source["delta"]:+g}.'))
        days = source.get('daily', [])
        if days:
            chart = Drawing(470, 105)
            maximum = max(1, *(day['observed'] for day in days))
            chart.add(Line(25,22,468,22,strokeColor=muted,strokeWidth=.4))
            for y, label in [(22,'0'),(86,str(maximum))]:
                chart.add(String(0,y,label,fontName='ERAUnicode',fontSize=7,fillColor=muted))
            step = 440 / len(days)
            for i, day in enumerate(days):
                chart.add(Rect(26+i*step,22,max(.4,step-1),day['observed']/maximum*62,
                    fillColor=accent if day['value'] is not None else colors.HexColor('#C5B7EB'),strokeColor=None))
            for x, day, anchor in [(25,days[0],'start'),(467,days[-1],'end')]:
                chart.add(String(x,6,day['date'],fontName='ERAUnicode',fontSize=7,fillColor=muted,textAnchor=anchor))
            story.append(chart)
        story.extend([p(source['explanation'], 'small'), p(source['history_note'], 'small'), Spacer(1, 4*mm)])
        story[source_start:] = [KeepTogether(story[source_start:])]
    section('05 / Команда и возможности', 'Обезличенные итоги. Департаменты и направления - текущий срез; достижения и заявки - за выбранные даты.')
    for key, title, fields in [('04','Департаменты',['Департамент','Участников','Активных целей','Выполнено целей']),
        ('05','Направления',['Направление','Департамент','Участников']),
        ('10','Заявки на возможности',['Статус заявки','Количество']),
        ('11','Достижения',['Тип достижения','Статус','Количество'])]:
        safe_rows = [row for row in records(key) if fields[0] in row]
        if not safe_rows:
            story.append(p(f'{title}: нет записей в доступном срезе.', 'small'))
            continue
        story.extend([CondPageBreak(24*mm), p(title, 'item')])
        for row in safe_rows:
            story.append(p(' · '.join(f'{label}: {text(row.get(label))}' for label in fields)))
    section('06 / Партнёрства и цели', 'Текущий организационный срез, не историческая оценка за месяц. Без контактных лиц и персональных целей.')
    contacts = getattr(analytics, 'contacts', []) if analytics else []
    story.append(p(f'Организаций в базе: {len(contacts)}.'))
    for contact in contacts:
        story.append(p(contact.organization_name))
    goals = [goal for goal in (getattr(analytics, 'goals', []) if analytics else [])
        if getattr(goal, 'scope_type', None) in ('organization', 'global', 'department', 'direction')]
    for goal in goals:
        story.extend([p(goal.title, 'item'), p(f'{goal.scope_name} · {goal.month} · выполнение {goal.current_value}/{goal.target_value}')])
    section('07 / Приоритеты руководителя', 'Сводка сигналов без персональных решений и свободных комментариев.')
    counts = {}
    for row in records('12'):
        status = row.get('Статус', 'не указан')
        counts[status] = counts.get(status, 0) + 1
    story.append(p('Сигналы за период: ' + (', '.join(f'{STATUS.get(key, key)} - {value}' for key, value in counts.items()) or 'нет записей')))
    risks = getattr(health, 'risks', [])
    recommendations = getattr(efficiency, 'recommendations', [])
    if risks or recommendations:
        story.append(p('Дополнительно - текущие агрегированные сигналы платформы. Их окна расчёта могут отличаться от периода отчёта.', 'small'))
    for risk in risks:
        story.append(p(risk))
    for recommendation in recommendations:
        story.extend([CondPageBreak(30*mm), p(recommendation.title, 'item'),
            p(recommendation.reason), p(recommendation.action)])
    section('Как читать отчёт', 'Методика и границы данных')
    story.append(p('Состав команды - на дату формирования. Проекты отобраны по созданию и подтверждённой работе; мероприятия - по календарной дате; задания - по созданию. '
        'Статусы текущие, не исторические. Telegram считает наблюдения с момента подключения: недоступная история не заменяется нулём. '
        'PDF и Excel используют один набор фактов. Индивидуальные данные и анкеты в этот отчёт не входят.', 'small'))
    wb.close()
    doc.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()
