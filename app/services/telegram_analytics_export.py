from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


def add_telegram_sheets(content: bytes, snapshot: dict) -> bytes:
    wb = load_workbook(BytesIO(content))
    ws = wb.create_sheet('Telegram · источники')
    ws.append(['Показатель', 'Значение за период', 'Учтено за период', 'Предыдущий период', 'Изменение',
               'Статус данных', 'Начало сбора', 'Непрерывное покрытие с', 'Период', 'Методика',
               'Учтено за 7 дней', 'Учтено за 30 дней', 'Учтено за 90 дней', 'Последнее поступление'])
    for src in snapshot['sources']:
        ws.append([src['label'], src['value'], src['observed_value'], src['previous_value'], src['delta'],
                   src['data_state'], src['started_at'], src['coverage_since'], f"{snapshot['start']}–{snapshot['end']}",
                   src['explanation'] + ' ' + src['history_note'],
                   *[next((item['observed'] for item in src.get('rolling_counts', []) if item['days'] == days), None) for days in (7,30,90)],
                   src.get('last_event_at')])
    days = wb.create_sheet('Telegram · по дням')
    days.append(['Источник', 'Дата', 'Значение', 'Учтено', 'Покрытие'])
    for src in snapshot['sources']:
        for row in src['daily']:
            days.append([src['label'], row['date'], row['value'], row['observed'],
                         'valid' if row['value'] is not None else 'insufficient'])
    for sheet in (ws, days):
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='632CFF')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 26
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
