from __future__ import annotations

from datetime import datetime, time, timedelta, timezone
from io import BytesIO
from html import escape
from zoneinfo import ZoneInfo

from sqlalchemy import func, select
from app.database.leadership_models import (
    WeeklyPulseArchive,
    WeeklyPulseCycle,
    LeadershipReportPulse,
)
from app.database.models import User, Task, TaskSubmission, Project, Event, EventRegistration, LeadershipReport
from app.utils.constants import TaskStatus


async def freeze_week(session, cycle: WeeklyPulseCycle) -> WeeklyPulseArchive:
    # Serialize competing scheduler/manual archive creation on PostgreSQL.
    await session.execute(
        select(WeeklyPulseCycle)
        .where(WeeklyPulseCycle.id == cycle.id)
        .with_for_update()
    )
    archived = await session.scalar(
        select(WeeklyPulseArchive).where(WeeklyPulseArchive.cycle_id == cycle.id)
    )
    if archived:
        return archived
    from app.services.weekly_pulse_cycle_service import get_schedule, sync_cycle_participants
    await sync_cycle_participants(session,cycle)
    schedule=await get_schedule(session)
    start = datetime.combine(cycle.date_from, time.min, tzinfo=ZoneInfo(schedule.timezone))
    end = datetime.combine(
        cycle.date_to + timedelta(days=1), time.min, tzinfo=ZoneInfo(schedule.timezone)
    )
    count = lambda model, *where: select(func.count()).select_from(model).where(*where)
    participants = int(
        await session.scalar(
            count(User, User.created_at < end, User.is_archived.is_(False), User.is_blocked.is_(False))
        )
        or 0
    )
    new_users = int(
        await session.scalar(
            count(
                User,
                User.created_at >= start,
                User.created_at < end,
                User.is_archived.is_(False),
            )
        )
        or 0
    )
    reports = list(
        (
            await session.scalars(
                select(LeadershipReport)
                .where(
                    LeadershipReport.pulse_cycle_id == cycle.id,
                    LeadershipReport.submitted_at.is_not(None),
                )
                .order_by(LeadershipReport.id)
            )
        ).all()
    )
    pulse_answers = {
        p.report_id: (p.answers_json or {}).get("submitted", {})
        for p in (
            await session.scalars(
                select(LeadershipReportPulse).where(
                    LeadershipReportPulse.report_id.in_([r.id for r in reports])
                )
            )
        ).all()
    }
    completed = list(
        (
            await session.scalars(
                select(Task)
                .where(
                    Task.status == TaskStatus.COMPLETED,
                    Task.id.in_(select(TaskSubmission.task_id).where(
                        TaskSubmission.status == "approved",TaskSubmission.updated_at >= start,TaskSubmission.updated_at < end)),
                )
                .order_by(Task.id)
            )
        ).all()
    )
    event_rows = list(
        (
            await session.scalars(
                select(Event)
                .where(
                    Event.event_date >= cycle.date_from,
                    Event.event_date <= cycle.date_to,
                    Event.status == "completed",
                )
                .order_by(Event.id)
            )
        ).all()
    )
    attended = list(
        (
            await session.scalars(
                select(EventRegistration.user_id)
                .join(Event, Event.id == EventRegistration.event_id)
                .where(
                    Event.event_date >= cycle.date_from,
                    Event.event_date <= cycle.date_to,
                    EventRegistration.status == "attended",
                )
            )
        ).all()
    )
    active = (
        set(attended)
        | {t.assignee_id for t in completed if t.assignee_id}
        | {r.owner_id for r in reports}
    )
    blocked = list(
        (
            await session.scalars(
                select(Task)
                .where(
                    Task.comment.is_not(None),
                    Task.comment != "",
                    Task.status.notin_([TaskStatus.COMPLETED, TaskStatus.CANCELLED]),
                )
                .order_by(Task.id)
            )
        ).all()
    )
    projects=list((await session.scalars(select(Project).where(Project.created_at < end))).all())
    previous=await session.scalar(select(WeeklyPulseArchive).join(WeeklyPulseCycle,WeeklyPulseCycle.id==WeeklyPulseArchive.cycle_id)
        .where(WeeklyPulseCycle.date_from < cycle.date_from).order_by(WeeklyPulseCycle.date_from.desc()).limit(1))
    snapshot = {
        "schema_version": 1,
        "period": f"{cycle.date_from:%d.%m.%Y} — {cycle.date_to:%d.%m.%Y}",
        "week": cycle.week_number,
        "year": cycle.date_from.isocalendar().year,
        "frozen_at": datetime.now(timezone.utc).isoformat(),
        "metrics": {
            "participants": participants,
            "new_participants": new_users,
            "active_participants": len(active),
            "submitted": len({r.owner_id for r in reports}),
            "eligible": cycle.eligible_count,
            "completed_tasks": len(completed),
            "events": len(event_rows),
            "attendance": len(attended),
            "active_projects": sum(p.status not in {"completed","cancelled","rejected","archived","draft"} for p in projects),
        },
        "active_definition": "Посещение мероприятия, завершённая назначенная задача или отправленный Пульс за период.",
        "reports": [
            {
                "id": r.id,
                "owner_id": r.owner_id,
                "result": r.main_result or "",
                "blocker": r.blocker_note or "",
                "priorities": r.next_priorities or [],
                "needs_help": r.needs_help,
                "answers": pulse_answers.get(r.id, {}),
            }
            for r in reports
        ],
        "tasks": [
            {"id": t.id, "title": t.title, "assignee_id": t.assignee_id}
            for t in completed
        ],
        "blockers": [
            {"id": t.id, "title": t.title, "text": t.comment} for t in blocked
        ],
        "events": [{"id": e.id, "title": e.title} for e in event_rows],
        "projects": [{"id":p.id,"title":p.title,"status":p.status} for p in projects],
        "previous_metrics": previous.snapshot.get("metrics") if previous else None,
        "opportunities": [{"report_id":r.id,"text":pulse_answers.get(r.id,{}).get("contacts","")} for r in reports if pulse_answers.get(r.id,{}).get("contacts")],
        "media_leads": [{"report_id":r.id,"text":pulse_answers.get(r.id,{}).get("media","")} for r in reports if pulse_answers.get(r.id,{}).get("media")],
        "warnings": [
            "Telegram: нет подтверждённых данных о просмотрах и реакциях за период.",
            "Сводка содержит ответы участников без автоматического подтверждения внешних фактов.",
            "Завершённые задачи учитываются по дате одобрения результата. Данные отражают доступное состояние на момент фиксации.",
        ],
    }
    import asyncio
    # Render before declaring the cycle complete; a failed renderer leaves a retryable cycle.
    docx_data = await asyncio.to_thread(render_docx,snapshot)
    pdf_data = await asyncio.to_thread(render_pdf,snapshot)
    archived = WeeklyPulseArchive(cycle_id=cycle.id, snapshot=snapshot,docx_data=docx_data,pdf_data=pdf_data)
    session.add(archived)
    cycle.status = "completed"
    cycle.report_generated_at = datetime.now(timezone.utc)
    cycle.submitted_count = snapshot["metrics"]["submitted"]
    await session.flush()
    return archived


def report_pages(snapshot):
    m = snapshot["metrics"]
    short = lambda value: " ".join(str(value).split())[:170]
    results = [
        f"Ответ #{r['id']}: {short(r['result'])}"
        for r in snapshot["reports"]
        if r["result"]
    ]
    tasks = [f"Задача #{t['id']}: {short(t['title'])}" for t in snapshot["tasks"]]
    blockers = [f"Задача #{b['id']}: {short(b['text'])}" for b in snapshot["blockers"]]
    blockers += [
        f"Ответ #{r['id']}: {short(r['blocker'])}"
        for r in snapshot["reports"]
        if r["blocker"]
    ]
    plans = [
        f"Ответ #{r['id']}: {short(p)}"
        for r in snapshot["reports"]
        for p in r["priorities"]
    ]

    def top(items, limit):
        return items[:limit] + (
            [f"Ещё {len(items) - limit}: полный состав сохранён в архиве ЭРА."]
            if len(items) > limit
            else []
        )

    return [
        [
            ("Пульс ЭРА", [snapshot["period"], "Внутренний отчёт для руководства"]),
            (
                "Показатели недели",
                [
                    f"Участников: {m['participants']} · Новых: {m['new_participants']}",
                    f"Активных: {m['active_participants']} · Ответили: {m['submitted']} / {m['eligible']}",
                    f"Завершённых задач: {m['completed_tasks']}",
                    f"Проведённых мероприятий: {m['events']} · Посещений: {m['attendance']}",
                ],
            ),
            (
                "Главное по ответам команды",
                top(results, 5) or ["Ответы с результатами не поступили."],
            ),
        ],
        [
            ("Люди и работа", [snapshot["period"]]),
            (
                "Завершённые задачи",
                top(tasks, 5) or ["Нет завершённых задач за период."],
            ),
            ("Требует решения", top(blockers, 5) or ["Блокеры не зафиксированы."]),
        ],
        [
            ("Следующая неделя", top(plans, 5) or ["Приоритеты не указаны."]),
            (
                "Источники и ограничения",
                snapshot["warnings"]
                + [
                    snapshot["active_definition"],
                    "Данные зафиксированы: " + snapshot["frozen_at"],
                    "Полные ответы и ссылки на исходные записи доступны администрации в архиве.",
                ],
            ),
        ],
    ]


def render_docx(snapshot) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "DejaVu Sans"
        doc.styles[name].font.color.rgb = RGBColor.from_string("303030")
    section.header.paragraphs[0].text = "ЭРА  /  ЕЖЕНЕДЕЛЬНЫЙ ОБЗОР"
    section.footer.paragraphs[
        0
    ].text = (
        f"ERA PULSE · {snapshot['year']} / W{snapshot['week']:02d} · Для руководства"
    )
    for idx, page in enumerate(report_pages(snapshot)):
        if idx:
            doc.add_page_break()
        for heading, lines in page:
            doc.add_heading(heading, level=1)
            for line in lines:
                doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(snapshot) -> bytes:
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4

    font = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    pdfmetrics.registerFont(TTFont("EraPulse", font))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "EraPulse"
    styles["Normal"].fontSize = 10
    styles["Normal"].leading = 14
    styles["Heading1"].fontSize = 18
    styles["Heading1"].leading = 22
    story = []
    for idx, page in enumerate(report_pages(snapshot)):
        if idx:
            story.append(PageBreak())
        for title, lines in page:
            story.append(Paragraph(escape(title), styles["Heading1"]))
            for line in lines:
                story += [Paragraph(escape(line), styles["Normal"]), Spacer(1, 8)]
    buf = BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, rightMargin=56, leftMargin=56, topMargin=50, bottomMargin=50
    ).build(story)
    return buf.getvalue()


async def archive_due_cycles(bot, settings, session_factory):
    async with session_factory() as session:
        cycles = list(
            (
                await session.scalars(
                    select(WeeklyPulseCycle).where(
                        WeeklyPulseCycle.closes_at <= datetime.now(timezone.utc),
                        WeeklyPulseCycle.status.notin_(["completed", "archived"]),
                    )
                )
            ).all()
        )
        for cycle in cycles:
            try:
                async with session.begin_nested():
                    await freeze_week(session, cycle)
                await session.commit()
            except Exception:
                import logging
                logging.getLogger(__name__).exception("Pulse archive failed cycle=%s",cycle.id)
                continue
