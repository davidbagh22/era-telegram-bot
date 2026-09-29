from __future__ import annotations

import io
import json
import logging
import os
from datetime import datetime, timezone
from typing import Any

from aiogram import F
from aiogram.enums import ChatType
from aiogram.types import BufferedInputFile, KeyboardButton, Message, ReplyKeyboardMarkup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.commission_bot import _kb
from app.commission_bot_community import (
    INTEREST_LABELS,
    STATUS_LABELS,
    SUPPORT_TOPICS,
    CommissionBotCommunity,
    _json_list,
)
from app.commission_bot_ultimate import _safe

log = logging.getLogger(__name__)

NAVIGATION_ACTIONS = {
    "➕ Создать",
    "✅ Согласование",
    "📋 Участники",
    "📋 Регистрации",
    "⚙️ Управление",
    "📅 События",
    "📅 Мероприятия",
    "🎟 Мои заявки",
    "🎟 Мои регистрации",
    "💬 Связаться",
    "💬 Обращения",
    "👤 Профиль",
    "🔔 Уведомления",
    "📊 Аналитика",
}


def should_interrupt_flow(state: str | None, onboarding_complete: bool, text: str) -> bool:
    """Main-menu actions intentionally leave an unfinished flow instead of trapping the user."""
    if not state or text not in NAVIGATION_ACTIONS:
        return False
    if state.startswith("onboard:") and not onboarding_complete:
        return False
    return True


def _percentage(part: int | float | None, whole: int | float | None) -> float:
    if not whole:
        return 0.0
    return round(float(part or 0) * 100.0 / float(whole), 1)


def _minutes_label(value: float | int | None) -> str:
    if value is None:
        return "—"
    minutes = max(float(value), 0.0)
    if minutes < 60:
        return f"{round(minutes)} мин"
    hours = minutes / 60
    if hours < 24:
        return f"{hours:.1f} ч"
    return f"{hours / 24:.1f} дн"


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(timezone.utc).replace(tzinfo=None)
        return value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _style_table_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        width = 10
        for cell in column_cells[:200]:
            value = cell.value
            if value is None:
                continue
            width = max(width, min(len(str(value)) + 2, 42))
            if isinstance(value, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_table(wb: Workbook, name: str, headers: list[str], rows: list[list[Any]]):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([_xlsx_value(value) for value in row])
    _style_table_sheet(ws)
    return ws


def build_analytics_workbook(snapshot: dict[str, Any]) -> bytes:
    """Build an admin-only XLSX report from a DB snapshot."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    summary = snapshot.get("summary", {})
    participants = int(summary.get("participants") or 0)
    profiles = int(summary.get("profiles") or 0)
    registrations = int(summary.get("registrations") or 0)
    attended = int(summary.get("attended") or 0)
    sent = int(summary.get("sent") or 0)
    failed = int(summary.get("failed") or 0)

    dashboard_rows = [
        ["Отчёт", "Аналитика Комиссии по взаимодействию с молодёжью"],
        ["Сформирован, UTC", datetime.now(timezone.utc).replace(tzinfo=None)],
        ["", ""],
        ["Участники", participants],
        ["Заполненные профили", profiles],
        ["Заполненность профилей, %", _percentage(profiles, participants)],
        ["Стран", int(summary.get("countries") or 0)],
        ["Подписаны на общие рассылки", int(summary.get("optins") or 0)],
        ["", ""],
        ["Публикации", int(summary.get("publications") or 0)],
        ["Опубликовано", int(summary.get("published") or 0)],
        ["Ближайшие события", int(summary.get("upcoming") or 0)],
        ["Регистрации", registrations],
        ["Уникальные участники событий", int(summary.get("unique_participants") or 0)],
        ["Посетили", attended],
        ["Посещаемость, %", _percentage(attended, registrations)],
        ["Лист ожидания", int(summary.get("waitlist") or 0)],
        ["", ""],
        ["Доставлено сообщений", sent],
        ["Ошибок доставки", failed],
        ["Успешная доставка, %", _percentage(sent, sent + failed)],
        ["Подключено чатов/каналов", int(summary.get("targets") or 0)],
        ["", ""],
        ["Обращения: ждут ответа", int(summary.get("support_open") or 0)],
        ["Обращения: отвечено", int(summary.get("support_answered") or 0)],
        ["Обращения: закрыто", int(summary.get("support_closed") or 0)],
        ["Среднее время первого ответа, мин", round(float(summary.get("avg_first_response_min") or 0), 1)],
    ]
    for row in dashboard_rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["B1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["B1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 40
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = "dd.mm.yyyy hh:mm"

    _append_table(
        wb,
        "Participants",
        [
            "Telegram ID", "ФИО", "Username", "Возраст", "Страна", "Город", "Статус",
            "Организация", "Телефон", "Email", "Соцсеть", "Интересы", "Рассылки",
            "Профиль заполнен", "Создан", "Обновлён",
        ],
        snapshot.get("participants", []),
    )
    _append_table(
        wb,
        "Events",
        [
            "ID", "Название", "Тип", "Статус", "Дата", "Формат", "Локация", "Мест",
            "Регистрации", "Пришли", "Лист ожидания", "Отменили", "Доставлено", "Ошибки",
            "Конверсия посещения, %",
        ],
        snapshot.get("events", []),
    )
    _append_table(
        wb,
        "Registrations",
        [
            "Event ID", "Событие", "Статус регистрации", "Дата регистрации", "Telegram ID", "ФИО",
            "Username", "Возраст", "Страна", "Город", "Статус участника", "Организация", "Телефон",
            "Email", "Соцсеть", "Интересы", "Ответы на вопросы",
        ],
        snapshot.get("registrations", []),
    )
    _append_table(
        wb,
        "Distribution",
        ["Event ID", "Событие", "Канал", "Всего попыток", "Отправлено", "Ошибки", "Успех, %"],
        snapshot.get("distribution", []),
    )
    _append_table(
        wb,
        "Support",
        [
            "Ticket ID", "Участник", "Страна", "Тема", "Статус", "Создано", "Последнее сообщение",
            "Ответственный", "Первый ответ, мин", "Сообщений",
        ],
        snapshot.get("support", []),
    )
    _append_table(
        wb,
        "Countries",
        ["Страна", "Участников", "Профили заполнены", "Рассылки включены", "Регистраций", "Пришли"],
        snapshot.get("countries", []),
    )
    _append_table(
        wb,
        "Interests",
        ["Интерес", "Участников"],
        snapshot.get("interests", []),
    )

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


class CommissionBotAnalytics(CommissionBotCommunity):
    def _quick_keyboard(self, role: str) -> ReplyKeyboardMarkup:
        if role in {"admin", "owner"}:
            return ReplyKeyboardMarkup(
                keyboard=[
                    [KeyboardButton(text="➕ Создать"), KeyboardButton(text="✅ Согласование")],
                    [KeyboardButton(text="📋 Участники"), KeyboardButton(text="📊 Аналитика")],
                    [KeyboardButton(text="📅 События"), KeyboardButton(text="🎟 Мои заявки")],
                    [KeyboardButton(text="💬 Обращения"), KeyboardButton(text="⚙️ Управление")],
                    [KeyboardButton(text="👤 Профиль"), KeyboardButton(text="❓ Помощь")],
                ],
                resize_keyboard=True,
                is_persistent=True,
                input_field_placeholder="Выберите действие",
            )
        return super()._quick_keyboard(role)

    async def _dispatch_ux_button(self, m: Message, user, text: str) -> None:
        st = await self.state_get(user["telegram_id"])
        state = str(st["state"]) if st else ""
        if st and should_interrupt_flow(state, bool(user["onboarding_complete"]), text):
            await self.state_clear(user["telegram_id"])
        elif st and state.startswith("onboard:") and not user["onboarding_complete"] and text in NAVIGATION_ACTIONS:
            await m.answer("Сначала выбери регион и страну — после этого откроется всё меню.")
            return
        if text == "📊 Аналитика":
            return await self._show_analytics(m.chat.id, user)
        await super()._dispatch_ux_button(m, user, text)

    async def _analytics_summary(self) -> dict[str, Any]:
        row = await self.fetchrow(
            """SELECT
            (SELECT COUNT(*) FROM users WHERE onboarding_complete) participants,
            (SELECT COUNT(*) FROM users WHERE participant_profile_complete) profiles,
            (SELECT COUNT(*) FROM users WHERE onboarding_complete AND telegram_opt_in) optins,
            (SELECT COUNT(DISTINCT country_code) FROM users WHERE onboarding_complete AND country_code IS NOT NULL) countries,
            (SELECT COUNT(*) FROM broadcasts) publications,
            (SELECT COUNT(*) FROM broadcasts WHERE status='published') published,
            (SELECT COUNT(*) FROM broadcasts WHERE status='published' AND event_at IS NOT NULL AND event_at>NOW()) upcoming,
            (SELECT COUNT(*) FROM registrations WHERE status IN ('registered','attended')) registrations,
            (SELECT COUNT(DISTINCT telegram_id) FROM registrations WHERE status IN ('registered','attended')) unique_participants,
            (SELECT COUNT(*) FROM registrations WHERE status='attended') attended,
            (SELECT COUNT(*) FROM registrations WHERE status='waitlist') waitlist,
            (SELECT COUNT(*) FROM deliveries WHERE status='sent') sent,
            (SELECT COUNT(*) FROM deliveries WHERE status='failed') failed,
            (SELECT COUNT(*) FROM targets WHERE status='approved') targets,
            (SELECT COUNT(*) FROM support_tickets WHERE status='open') support_open,
            (SELECT COUNT(*) FROM support_tickets WHERE status='answered') support_answered,
            (SELECT COUNT(*) FROM support_tickets WHERE status='closed') support_closed,
            (SELECT AVG(EXTRACT(EPOCH FROM (first_staff.first_response_at-t.created_at))/60.0)
               FROM support_tickets t
               JOIN LATERAL (
                 SELECT MIN(created_at) first_response_at FROM support_messages
                 WHERE ticket_id=t.id AND sender_role<>'viewer'
               ) first_staff ON first_staff.first_response_at IS NOT NULL
            ) avg_first_response_min"""
        )
        return dict(row or {})

    async def _show_analytics(self, chat_id: int, user) -> None:
        if user["role"] not in {"owner", "admin"}:
            return await self.bot.send_message(chat_id, "Аналитика доступна администратору и владельцу.")
        summary = await self._analytics_summary()
        countries = await self.fetch(
            """SELECT COALESCE(country_name,country_code,'Не указана') country,COUNT(*) participants
            FROM users WHERE onboarding_complete
            GROUP BY COALESCE(country_name,country_code,'Не указана')
            ORDER BY participants DESC,country LIMIT 8"""
        )
        events = await self.fetch(
            """SELECT b.id,b.title,
            COUNT(r.id) FILTER (WHERE r.status IN ('registered','attended')) registrations,
            COUNT(r.id) FILTER (WHERE r.status='attended') attended
            FROM broadcasts b LEFT JOIN registrations r ON r.broadcast_id=b.id
            WHERE b.status IN ('published','approved','scheduled')
            GROUP BY b.id,b.title
            ORDER BY registrations DESC,b.id DESC LIMIT 6"""
        )
        interests = await self.fetch(
            """SELECT value interest,COUNT(*) participants
            FROM users u CROSS JOIN LATERAL jsonb_array_elements_text(u.interests) value
            WHERE u.onboarding_complete
            GROUP BY value ORDER BY participants DESC,value LIMIT 6"""
        )

        participants = int(summary.get("participants") or 0)
        profiles = int(summary.get("profiles") or 0)
        registrations = int(summary.get("registrations") or 0)
        attended = int(summary.get("attended") or 0)
        sent = int(summary.get("sent") or 0)
        failed = int(summary.get("failed") or 0)
        text = (
            "<b>📊 Аналитика</b>\n\n"
            "<b>Аудитория</b>\n"
            f"👥 Участников: <b>{participants}</b>\n"
            f"🧩 Заполненных профилей: <b>{profiles}</b> ({_percentage(profiles, participants)}%)\n"
            f"🌍 Стран: <b>{int(summary.get('countries') or 0)}</b>\n"
            f"🔔 Получают общие рассылки: <b>{int(summary.get('optins') or 0)}</b>\n\n"
            "<b>Мероприятия</b>\n"
            f"📣 Публикаций: <b>{int(summary.get('publications') or 0)}</b> · опубликовано {int(summary.get('published') or 0)}\n"
            f"📅 Ближайших событий: <b>{int(summary.get('upcoming') or 0)}</b>\n"
            f"🎟 Регистраций: <b>{registrations}</b> · уникальных участников {int(summary.get('unique_participants') or 0)}\n"
            f"🚪 Пришли: <b>{attended}</b> · посещаемость {_percentage(attended, registrations)}%\n"
            f"⏳ В листах ожидания: <b>{int(summary.get('waitlist') or 0)}</b>\n\n"
            "<b>Распространение</b>\n"
            f"✅ Доставлено: <b>{sent}</b> · ошибок {failed} · успех {_percentage(sent, sent + failed)}%\n"
            f"📣 Подключено чатов/каналов: <b>{int(summary.get('targets') or 0)}</b>\n\n"
            "<b>Обращения</b>\n"
            f"🔴 Ждут ответа: <b>{int(summary.get('support_open') or 0)}</b>\n"
            f"🟢 Отвечено: {int(summary.get('support_answered') or 0)} · закрыто {int(summary.get('support_closed') or 0)}\n"
            f"⏱ Средний первый ответ: <b>{_minutes_label(summary.get('avg_first_response_min'))}</b>"
        )
        if countries:
            text += "\n\n<b>Топ стран</b>\n" + "\n".join(
                f"{i}. {_safe(row['country'])} — {row['participants']}" for i, row in enumerate(countries, 1)
            )
        if events:
            text += "\n\n<b>Топ событий по регистрациям</b>\n" + "\n".join(
                f"{i}. {_safe(row['title'])} — {row['registrations']}" for i, row in enumerate(events, 1)
            )
        if interests:
            text += "\n\n<b>Интересы аудитории</b>\n" + " · ".join(
                f"{_safe(INTEREST_LABELS.get(row['interest'], row['interest']))}: {row['participants']}" for row in interests
            )
        await self.bot.send_message(
            chat_id,
            text,
            reply_markup=_kb([
                [("📥 Выгрузить Excel", "analytics:excel")],
                [("🔄 Обновить", "ux:analytics"), ("⬅️ Меню", "menu")],
            ]),
        )

    async def _analytics_export_snapshot(self) -> dict[str, Any]:
        summary = await self._analytics_summary()
        participant_rows = await self.fetch(
            """SELECT telegram_id,registration_name,first_name,last_name,username,age,country_name,city,
            participant_status,organization,phone,email,social_url,interests,telegram_opt_in,
            participant_profile_complete,created_at,updated_at
            FROM users WHERE onboarding_complete ORDER BY country_name NULLS LAST,registration_name NULLS LAST,telegram_id"""
        )
        participants = []
        for row in participant_rows:
            name = row["registration_name"] or " ".join(v for v in [row["first_name"], row["last_name"]] if v).strip()
            participants.append([
                row["telegram_id"], name, "@" + row["username"] if row["username"] else None, row["age"], row["country_name"],
                row["city"], STATUS_LABELS.get(row["participant_status"], row["participant_status"]), row["organization"], row["phone"],
                row["email"], row["social_url"], ", ".join(INTEREST_LABELS.get(v, v) for v in _json_list(row["interests"])),
                "Да" if row["telegram_opt_in"] else "Нет", "Да" if row["participant_profile_complete"] else "Нет",
                row["created_at"], row["updated_at"],
            ])

        event_rows = await self.fetch(
            """SELECT b.id,b.title,b.content_type,b.status,b.event_at,b.event_format,b.event_location,b.capacity,
            COUNT(r.id) FILTER (WHERE r.status IN ('registered','attended')) registrations,
            COUNT(r.id) FILTER (WHERE r.status='attended') attended,
            COUNT(r.id) FILTER (WHERE r.status='waitlist') waitlist,
            COUNT(r.id) FILTER (WHERE r.status='cancelled') cancelled,
            (SELECT COUNT(*) FROM deliveries d WHERE d.broadcast_id=b.id AND d.status='sent') sent,
            (SELECT COUNT(*) FROM deliveries d WHERE d.broadcast_id=b.id AND d.status='failed') failed
            FROM broadcasts b LEFT JOIN registrations r ON r.broadcast_id=b.id
            GROUP BY b.id ORDER BY b.id DESC"""
        )
        events = [[
            row["id"], row["title"], row["content_type"], row["status"], row["event_at"], row["event_format"], row["event_location"],
            row["capacity"], row["registrations"], row["attended"], row["waitlist"], row["cancelled"], row["sent"], row["failed"],
            _percentage(row["attended"], row["registrations"]),
        ] for row in event_rows]

        registration_rows = await self.fetch(
            """SELECT r.broadcast_id,b.title,r.status,r.registered_at,r.answers,u.telegram_id,u.registration_name,u.first_name,u.last_name,
            u.username,u.age,u.country_name,u.city,u.participant_status,u.organization,u.phone,u.email,u.social_url,u.interests
            FROM registrations r JOIN broadcasts b ON b.id=r.broadcast_id JOIN users u ON u.telegram_id=r.telegram_id
            ORDER BY r.registered_at DESC"""
        )
        registrations = []
        for row in registration_rows:
            name = row["registration_name"] or " ".join(v for v in [row["first_name"], row["last_name"]] if v).strip()
            registrations.append([
                row["broadcast_id"], row["title"], row["status"], row["registered_at"], row["telegram_id"], name,
                "@" + row["username"] if row["username"] else None, row["age"], row["country_name"], row["city"],
                STATUS_LABELS.get(row["participant_status"], row["participant_status"]), row["organization"], row["phone"], row["email"],
                row["social_url"], ", ".join(INTEREST_LABELS.get(v, v) for v in _json_list(row["interests"])), dict(row["answers"] or {}),
            ])

        distribution_rows = await self.fetch(
            """SELECT d.broadcast_id,b.title,d.channel,COUNT(*) attempts,
            COUNT(*) FILTER (WHERE d.status='sent') sent,COUNT(*) FILTER (WHERE d.status='failed') failed
            FROM deliveries d JOIN broadcasts b ON b.id=d.broadcast_id
            GROUP BY d.broadcast_id,b.title,d.channel ORDER BY d.broadcast_id DESC,d.channel"""
        )
        distribution = [[
            row["broadcast_id"], row["title"], row["channel"], row["attempts"], row["sent"], row["failed"],
            _percentage(row["sent"], row["sent"] + row["failed"]),
        ] for row in distribution_rows]

        support_rows = await self.fetch(
            """SELECT t.id,t.user_id,t.topic,t.status,t.created_at,t.last_message_at,u.registration_name,u.first_name,u.last_name,u.username,u.country_name,
            staff.registration_name staff_name,staff.first_name staff_first,staff.last_name staff_last,staff.username staff_username,
            (SELECT MIN(sm.created_at) FROM support_messages sm WHERE sm.ticket_id=t.id AND sm.sender_role<>'viewer') first_response_at,
            (SELECT COUNT(*) FROM support_messages sm WHERE sm.ticket_id=t.id) messages
            FROM support_tickets t JOIN users u ON u.telegram_id=t.user_id
            LEFT JOIN users staff ON staff.telegram_id=t.assigned_to ORDER BY t.id DESC"""
        )
        support = []
        for row in support_rows:
            participant = row["registration_name"] or " ".join(v for v in [row["first_name"], row["last_name"]] if v).strip() or row["username"]
            staff_name = row["staff_name"] or " ".join(v for v in [row["staff_first"], row["staff_last"]] if v).strip() or row["staff_username"]
            first_response = None
            if row["first_response_at"]:
                first_response = round((row["first_response_at"] - row["created_at"]).total_seconds() / 60.0, 1)
            support.append([
                row["id"], participant, row["country_name"], SUPPORT_TOPICS.get(row["topic"], row["topic"]), row["status"], row["created_at"], row["last_message_at"],
                staff_name, first_response, row["messages"],
            ])

        country_rows = await self.fetch(
            """SELECT COALESCE(u.country_name,u.country_code,'Не указана') country,
            COUNT(DISTINCT u.telegram_id) participants,
            COUNT(DISTINCT u.telegram_id) FILTER (WHERE u.participant_profile_complete) profiles,
            COUNT(DISTINCT u.telegram_id) FILTER (WHERE u.telegram_opt_in) optins,
            COUNT(r.id) FILTER (WHERE r.status IN ('registered','attended')) registrations,
            COUNT(r.id) FILTER (WHERE r.status='attended') attended
            FROM users u LEFT JOIN registrations r ON r.telegram_id=u.telegram_id
            WHERE u.onboarding_complete
            GROUP BY COALESCE(u.country_name,u.country_code,'Не указана') ORDER BY participants DESC,country"""
        )
        countries = [[row["country"], row["participants"], row["profiles"], row["optins"], row["registrations"], row["attended"]] for row in country_rows]

        interest_rows = await self.fetch(
            """SELECT value interest,COUNT(*) participants
            FROM users u CROSS JOIN LATERAL jsonb_array_elements_text(u.interests) value
            WHERE u.onboarding_complete GROUP BY value ORDER BY participants DESC,value"""
        )
        interests = [[INTEREST_LABELS.get(row["interest"], row["interest"]), row["participants"]] for row in interest_rows]

        return {
            "summary": summary,
            "participants": participants,
            "events": events,
            "registrations": registrations,
            "distribution": distribution,
            "support": support,
            "countries": countries,
            "interests": interests,
        }

    def _register_handlers(self):
        r = self.router

        @r.callback_query(F.data == "analytics:excel")
        async def analytics_excel(c):
            await c.answer("Готовлю Excel…")
            user = await self.ensure_user(c.from_user)
            if user["role"] not in {"owner", "admin"}:
                return await c.message.answer("Выгрузка доступна администратору и владельцу.")
            try:
                snapshot = await self._analytics_export_snapshot()
                data = build_analytics_workbook(snapshot)
                stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
                await self.bot.send_document(
                    c.message.chat.id,
                    BufferedInputFile(data, filename=f"commission_analytics_{stamp}.xlsx"),
                    caption="📊 Полная аналитика: участники, события, регистрации, рассылки, обращения, страны и интересы.",
                )
                await self.audit(user["telegram_id"], "analytics_exported", "analytics", None, {"format": "xlsx"})
            except Exception:
                log.exception("Could not export Commission analytics workbook")
                await c.message.answer("Не получилось сформировать Excel. Ошибка записана в журнал, попробуйте ещё раз через минуту.")

        @r.message(F.chat.type == ChatType.PRIVATE, F.text == "📊 Аналитика")
        async def analytics_quick(m: Message):
            user = await self.ensure_user(m.from_user)
            await self._dispatch_ux_button(m, user, "📊 Аналитика")

        super()._register_handlers()


async def run_commission_bot_analytics(database_url: str) -> None:
    token = os.getenv("COMMISSION_BOT_TOKEN", "").strip()
    if not token:
        log.info("COMMISSION_BOT_TOKEN not set; commission bot disabled")
        return
    bootstrap = os.getenv("COMMISSION_BOOTSTRAP_CODE", "").strip()
    app = CommissionBotAnalytics(database_url, token, bootstrap)
    await app.run()
